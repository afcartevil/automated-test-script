#!/usr/bin/env python3
"""Export patch_analyzer JSON to an Excel workbook (.xlsx)."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _cell(ws, row: int, col: int, value: Any) -> None:
    if value is None:
        value = ""
    elif isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    ws.cell(row=row, column=col, value=value)


def write_summary_sheet(ws, data: Dict[str, Any]) -> None:
    ws.append(["字段", "值"])
    header = ws[1]
    for h in header:
        h.font = Font(bold=True)
    r = 2
    rows = [
        ("repo", data.get("repo", "")),
        ("from_tag", data.get("from_tag", "")),
        ("to_tag", data.get("to_tag", "")),
        ("total", data.get("total", "")),
        ("analyzed", data.get("analyzed", "")),
        ("failed", data.get("failed", "")),
    ]
    for k, v in rows:
        _cell(ws, r, 1, k)
        _cell(ws, r, 2, v)
        r += 1
    s = data.get("summary") or {}
    _cell(ws, r, 1, "summary (JSON)")
    _cell(ws, r, 2, json.dumps(s, ensure_ascii=False, indent=2))
    r += 1
    fails = data.get("failed_items") or []
    if fails:
        _cell(ws, r, 1, "failed_items (JSON)")
        _cell(ws, r, 2, json.dumps(fails, ensure_ascii=False, indent=2))


def _join_labels(tags: Any, key: str = "label") -> str:
    if not isinstance(tags, list):
        return ""
    out: List[str] = []
    for t in tags:
        if isinstance(t, dict):
            lab = t.get(key)
            if lab:
                out.append(str(lab))
        elif t:
            out.append(str(t))
    return "；".join(out)


def _dims_row(p: Dict[str, Any]) -> Dict[str, Any]:
    d = p.get("dimensions") or {}
    return {
        "dim_正确性与完整性": d.get("correctness_completeness", ""),
        "dim_风险与安全性": d.get("risk_safety", ""),
        "dim_Diff可维护性": d.get("diff_maintainability", ""),
        "dim_Commit元信息": d.get("commit_metadata_quality", ""),
    }


def _why_scored(p: Dict[str, Any]) -> str:
    w = p.get("why_scored")
    if isinstance(w, list):
        return "\n".join(str(x) for x in w)
    return str(w or "")


def _test_suggestions(p: Dict[str, Any]) -> str:
    ts = p.get("test_suggestions")
    if not isinstance(ts, dict):
        return str(ts or "")
    lines = []
    for k, v in ts.items():
        lines.append(f"{k}: {v}")
    return "\n".join(lines)


def _tag_labels_string(p: Dict[str, Any], simple_key: str, rich_key: str) -> str:
    raw = p.get(simple_key)
    if isinstance(raw, list) and raw:
        if all(isinstance(x, str) for x in raw):
            return "；".join(raw)
    return _join_labels(p.get(rich_key))


# 与 patch_analyzer.merge_patch_tags_with_intent 一致：intent_* 为 fix bug / 新需求
_INTENT_TAG_KEYS_PREFIX = "intent_"
_INTENT_LABELS = frozenset({"fix bug", "新需求"})


def _split_intent_and_module_tags(p: Dict[str, Any]) -> tuple[str, str]:
    """拆成：意图列（fix bug / 新需求）、模块/领域列（文件系统、驱动框架等）。"""
    tags = p.get("patch_tags")
    if isinstance(tags, list) and tags:
        intent_parts: List[str] = []
        module_parts: List[str] = []
        for t in tags:
            if not isinstance(t, dict):
                continue
            key = str(t.get("key") or "")
            lab = str(t.get("label") or "").strip()
            if not lab:
                continue
            is_intent = key.startswith(_INTENT_TAG_KEYS_PREFIX) or lab in _INTENT_LABELS
            if is_intent:
                intent_parts.append(lab)
            else:
                module_parts.append(lab)
        return "；".join(intent_parts), "；".join(module_parts)

    # 无结构化 patch_tags 时，仅按 patch_tag_labels 文本启发式拆分
    raw = p.get("patch_tag_labels")
    if isinstance(raw, list) and raw:
        intent = [str(x) for x in raw if str(x).strip() in _INTENT_LABELS]
        module = [str(x) for x in raw if str(x).strip() not in _INTENT_LABELS]
        return "；".join(intent), "；".join(module)
    return "", ""


def write_patches_sheet(ws, patches: List[Dict[str, Any]]) -> None:
    dim_keys = ["dim_正确性与完整性", "dim_风险与安全性", "dim_Diff可维护性", "dim_Commit元信息"]
    headers = [
        "commit",
        "title",
        "score",
        "recommendation",
        "blocked",
        "trusted_integration_merge",
        "active_profile",
        "key_risk",
        "意图标签",
        "模块标签",
        "platform_tags",
        "files_changed",
        "insertions",
        "deletions",
        *dim_keys,
        "files",
        "problem_solved",
        "change_summary",
        "why_scored",
        "test_suggestions",
    ]

    ws.append(headers)
    for h in ws[1]:
        h.font = Font(bold=True)

    for p in patches:
        dims = _dims_row(p)
        fl = p.get("files") or []
        files_str = "; ".join(str(x) for x in fl) if isinstance(fl, list) else str(fl)
        intent_str, module_str = _split_intent_and_module_tags(p)
        row = [
            p.get("commit", ""),
            p.get("title", ""),
            p.get("score", ""),
            p.get("recommendation", ""),
            p.get("blocked", ""),
            p.get("trusted_integration_merge", ""),
            p.get("active_profile", ""),
            p.get("key_risk", ""),
            intent_str,
            module_str,
            _tag_labels_string(p, "platform_tag_labels", "platform_tags"),
            p.get("files_changed", ""),
            p.get("insertions", ""),
            p.get("deletions", ""),
            dims.get("dim_正确性与完整性", ""),
            dims.get("dim_风险与安全性", ""),
            dims.get("dim_Diff可维护性", ""),
            dims.get("dim_Commit元信息", ""),
            files_str,
            p.get("problem_solved", ""),
            p.get("change_summary", ""),
            _why_scored(p),
            _test_suggestions(p),
        ]
        ws.append(row)

    for col_idx, name in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        width = min(52, max(10, min(len(name) + 2, 24) + (8 if "why" in name or name == "problem_solved" else 0)))
        ws.column_dimensions[letter].width = width


def main() -> int:
    parser = argparse.ArgumentParser(description="Export patch_analyzer JSON to Excel")
    parser.add_argument("--input", required=True, help="path to patch analyzer output JSON")
    parser.add_argument("--output", help="output .xlsx path (default: same basename as input)")
    args = parser.parse_args()

    in_path = Path(args.input).resolve()
    if not in_path.is_file():
        print(f"错误: 找不到输入文件: {in_path}", flush=True)
        return 2

    out_path = Path(args.output).resolve() if args.output else in_path.with_suffix(".xlsx")

    with in_path.open("r", encoding="utf-8") as f:
        data = json.load(f)

    patches = data.get("patches") or []
    if not isinstance(patches, list):
        print("错误: JSON 中无 patches 列表", flush=True)
        return 2

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "summary"
    write_summary_sheet(ws0, data)

    ws1 = wb.create_sheet("patches")
    write_patches_sheet(ws1, patches)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"已写入: {out_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
