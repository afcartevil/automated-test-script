"""Export a Feishu-style GKI upgrade checklist workbook from analysis results."""
from __future__ import annotations

import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import quote

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from web_app.analyzer_service import _parse_month_subtag, is_key_risk_none
from web_app.app_paths import checklist_config_path
from web_app.database import get_patches_by_shas, get_task, get_task_patches
from web_app.dev_module_classifier import get_dev_module_names, sort_modules_by_definition

# 与飞书模板 wiki/CZL7wHo8giJujkk07xWcHLtgntb、ChlPwxSWBidR9kkLbMYcdALRnuh 对齐，并按验收反馈调整：
# - 表格只放关键信息；细节见「模块详情」链接
# - 「重点 Review」按 AI 信号做 P0/P1/P2 分级，突出必看变更
# - 「测试用例」填 AI 分析结论
# - 「模块 Review 意见」留空给 Reviewer/Owner 填写 review 结果与人工测试用例
OVERVIEW_HEADERS = [
    "",  # A：与 B 合并为 Tag 列
    "GKI Tag / Changes",
    "提交版本信息",
    "二级部门",
    "开发领域&Owner",
    "模块详情",
    "审核人",
    "review+1",
    "变更摘要",
    "重点 Review",
    "是否有负面影响",
    "风险项描述",
    "是否必须合入",
    "测试用例",
    "模块 Review 意见",
]

BASE_TEST_HEADERS = [
    "二级部门",
    "开发领域",
    "模块 Owner",
    "🚨 GKI升级时：是否必须纳入【基础测试】？(填：是/否)",
    "📝 推荐的基础测试场景/核心用例 (若选是，请简述)",
    "🤖 现有自动化覆盖情况 (全量/部分/无)",
    "请确认如下测试输出的用例与E列的“基础测试场景”是否匹配",
]

_HEADER_PRINCIPLES = {
    "代码审核原则": (
        "基本要求：选中 r1 后除了 review r1 变更，还需要 review r2→rN 变更，"
        "防止 r2→rN 变更存在修复 r1 严重故障的 patch。\n"
        "🔵 r1 (领域兜底)：\n"
        "领域负责制：r1 变更代码量巨大，基本上都要 1~2 千个文件，没法按照责任田对号入座。"
        "由修改点影响到的最相关业务领域负责。\n"
        "风险兜底：直接执行“领域兜底原则”。一旦测试发现 r1 引入的异常，依据代码功能属性"
        "（如内存、调度、驱动）直接由对应模块 Owner 负责修复，不得以“未 Review”为由推诿。\n"
        "🔴 r2 → rN (增量必检)：\n"
        "强制 Review：必须逐笔分配到负责人仔细 review。\n"
        "防回退：重点确认 r2-rN 是否存在修复严重 bug 的变更；若存在需要反馈给 GKI 升级负责人，"
        "评估是否升级到 rN。"
    ),
    "测试与验证策略": (
        "验证模式：“基础测试” + “研发修改建议” 叠加模式。\n"
        "执行逻辑：测试团队默认执行 GKI 升级保底的“基础测试”用例；"
        "各领域研发在 Review 代码后，需判断是否有超出基础用例覆盖范围的特殊风险。\n"
        "如果有，必须在下方表格提出“专项测试建议”；如无，则默认由基础测试兜底。"
    ),
    "质量控制": (
        "1+N 编译和发布 OTA 版本：\n"
        "1）必须在首选项目 OTA 释放给粉丝和用户 1 个月后，稳定性 APR 达标（不增长）"
        "才能陆续开放其它项目的编译、测试、外发。\n"
        "2）首选项目 GKI 升维团队负责每周监控系统 APR 指标。\n"
        "红线：只有首推项目 APR 达标，且上层业务给出“非内核相关异常”的结论后，"
        "才能开放其他项目的 OTA 测试与推送。"
    ),
    "质量定责": (
        "领域兜底原则 (归属认定)：\n"
        "代码 Review 或测试中发现的新问题，依据“代码功能属性”定责。\n"
        "判定逻辑：引入的问题代码与哪个业务领域最相关（如涉及内存归内核内存，"
        "进程相关归内核调度，fs 框架问题归存储等），即由该领域负责。\n"
        "无主地带：即使该模块未明确划分“责任田”，相关领域必须背责，严禁推诿。\n"
        "共担机制：如果影响共分支的外研和转维团队，外研和转维团队需要一起背责。"
    ),
}

# 配色与飞书模板 GKI升级专项checklist 保持一致（wiki/CZL7wHo8giJujkk07xWcHLtgntb）
_COLOR_HEADER_BG = "#049FD7"         # 列标题行背景（青蓝）
_COLOR_HEADER_FG = "#000000"         # 列标题行文字
_COLOR_BASE_HEADER_FG = "#333333"    # 基础测试需求表头文字
_COLOR_BASE_HEADER_ACCENT = "#FFF258"  # 基础测试表头最后一列强调色
_COLOR_BORDER = "#D0D0D0"

_FILL_HEADER = PatternFill("solid", fgColor="049FD7")
_FILL_BASE_HEADER_ACCENT = PatternFill("solid", fgColor="FFF258")
_FONT_HEADER = Font(bold=True, color="000000", size=10, name="微软雅黑")
_FONT_BASE_HEADER = Font(bold=False, color="333333", size=11, name="微软雅黑")
_FONT_LABEL = Font(bold=True, size=10, name="微软雅黑")
_FONT_BODY = Font(size=10, name="微软雅黑")
_ALIGN_WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
_ALIGN_META_LABEL = Alignment(wrap_text=True, vertical="center", horizontal="center")
_ALIGN_META_VALUE = Alignment(wrap_text=True, vertical="center", horizontal="left")
_ALIGN_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_NONE_RISK = frozenset({"无", "无风险", "none", "n/a", "na", "-", "—", "n.a."})
_TEXT_LIMIT = 1200
# 飞书单元格输出上限：表格只突出关键信息，细节放模块详情页
_CLIP_MODIFICATION = 520
_CLIP_FOCUS = 900
_CLIP_TEST = 420
_CLIP_RISK = 360
# 送入 LLM 打磨的草稿上限
_DRAFT_MODIFICATION = 1800
_DRAFT_FOCUS = 2200

# 列宽：A:B=Tag，其后提交/部门/Owner/模块详情/…/测试用例/模块 Review 意见
_COL_WIDTHS_OVERVIEW = {
    1: 22, 2: 22, 3: 28, 4: 14, 5: 26, 6: 12, 7: 12, 8: 16,
    9: 28, 10: 34, 11: 12, 12: 22, 13: 12, 14: 28, 15: 28,
}
_COL_WIDTHS_BASE = {1: 16, 2: 22, 3: 18, 4: 28, 5: 48, 6: 18, 7: 28}
_COL_PIXELS_OVERVIEW = {
    1: 170, 2: 170, 3: 210, 4: 120, 5: 200, 6: 100, 7: 100, 8: 120,
    9: 220, 10: 280, 11: 100, 12: 180, 13: 100, 14: 220, 15: 220,
}
_COL_PIXELS_BASE = {1: 150, 2: 170, 3: 120, 4: 220, 5: 380, 6: 150, 7: 220}
_OVERVIEW_WRAP = {1: 20, 3: 26, 9: 28, 10: 32, 12: 24, 14: 28, 15: 28}
_BASE_WRAP = {5: 42}
# 模块详情 / 审核人 / review+1 / 是否有负面影响 / 是否必须合入
_OVERVIEW_CENTER_COLS = [6, 7, 8, 11, 13]
_COL_TAG = 1
_COL_TAG_SPAN = 2
_COL_SUBMIT = 3
_COL_DEPT = 4
_COL_OWNER = 5
_COL_MODULE_LINK = 6
_COL_REVIEWER = 7
_COL_REVIEW_PLUS = 8
_COL_MODIFICATION = 9
_COL_FOCUS_REVIEW = 10
_COL_NEGATIVE = 11
_COL_RISK = 12
_COL_MUST_MERGE = 13
_COL_TEST = 14
_COL_OWNER_COMMENT = 15

# 与飞书模板「版本验收总览」下拉选项/颜色一致
_REVIEW_PLUS_OPTIONS = ["已完成 review+1", "不涉及", "已完成 review"]
_REVIEW_PLUS_COLORS = ["#34c724", "#fed4a4", "#b1e8fc"]
_YES_NO_NA_OPTIONS = ["是", "否", "NA"]
_YES_NO_NA_COLORS = ["#bacefd", "#fed4a4", "#b1e8fc"]
_BASE_CENTER_COLS = [4, 6, 7]


@dataclass
class ModuleOwner:
    department: str = ""
    lead: str = ""
    owner: str = ""
    reviewer: str = ""


@dataclass
class ChecklistRow:
    module_name: str = ""
    tag_text: str = ""
    submit_info: str = ""
    department: str = ""
    domain_owner: str = ""
    reviewer: str = ""
    review_plus: str = ""
    modification: str = ""
    focus_review: str = ""
    negative_impact: str = ""
    risk_description: str = ""
    must_merge: str = ""
    test_suggestions: str = ""
    owner_comment: str = ""  # 预留给 Reviewer/Owner 填写 review 结果与人工测试用例
    skipped: bool = False
    # overall = 总包；integration = 集成维护固定 check；subtag = 相邻增量涉及模块（不多列）
    block_kind: str = "overall"
    source_patches: List[Dict[str, Any]] = field(default_factory=list, repr=False)
    values: List[str] = field(default_factory=list)

    def to_values(self) -> List[str]:
        # A:B 合并为 Tag（内容写在 A，B 留空），与飞书模板一致
        # 「模块详情」列在导出时再填链接；
        # 「模块 Review 意见」始终留空给人填
        return [
            self.tag_text,
            "",
            self.submit_info,
            self.department,
            self.domain_owner,
            "",
            self.reviewer,
            self.review_plus,
            self.modification,
            self.focus_review,
            self.negative_impact,
            self.risk_description,
            self.must_merge,
            self.test_suggestions,
            "",  # 模块 Review 意见：人工填写
        ]


def build_checklist_xlsx(task_id: str) -> Tuple[bytes, str]:
    """Build the checklist workbook. Returns (xlsx_bytes, download_filename)."""
    task = get_task(task_id)
    if not task:
        raise FileNotFoundError("任务不存在")

    patches = [_hydrate_patch(row) for row in get_task_patches(task_id)]
    if not patches and task.get("type") != "range":
        raise ValueError("暂无分析结果，无法导出验收表格")

    wb = Workbook()
    ws = wb.active
    ws.title = "版本验收总览"
    _write_overview_sheet(ws, task, patches)

    ws2 = wb.create_sheet("基础测试需求")
    _write_base_test_sheet(ws2, task, patches)

    buf = io.BytesIO()
    wb.save(buf)
    filename = _checklist_filename(task)
    return buf.getvalue(), filename


def checklist_content_disposition(filename: str) -> str:
    ascii_name = "gki_upgrade_checklist.xlsx"
    encoded = quote(filename)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded}"


@dataclass
class SheetSpec:
    title: str
    values: List[List[Any]]
    merges: List[Tuple[int, int, int, int]] = field(default_factory=list)
    frozen_rows: int = 0
    col_widths: Dict[int, int] = field(default_factory=dict)
    col_pixels: Dict[int, int] = field(default_factory=dict)
    header_row: int = 1
    skip_rows: List[int] = field(default_factory=list)
    tag_blocks: List[Tuple[int, int, str]] = field(default_factory=list)
    center_cols: List[int] = field(default_factory=list)
    row_heights: Dict[int, int] = field(default_factory=dict)


def build_checklist_sheets(
    task_id: str,
    *,
    module_report_urls: Optional[Dict[str, str]] = None,
    subtag_report_urls: Optional[Dict[str, str]] = None,
    module_report_zip_url: str = "",
    module_report_folder_url: str = "",
) -> Tuple[str, List[SheetSpec]]:
    """Build in-memory sheets used by both xlsx and Feishu exporters."""
    task = get_task(task_id)
    if not task:
        raise FileNotFoundError("任务不存在")
    patches = [_hydrate_patch(row) for row in get_task_patches(task_id)]
    if not patches and task.get("type") != "range":
        raise ValueError("暂无分析结果，无法导出验收表格")
    title = _checklist_filename(task).replace(".xlsx", "")
    return title, [
        _overview_sheet_spec(
            task,
            patches,
            module_report_urls=module_report_urls or {},
            subtag_report_urls=subtag_report_urls or {},
            module_report_zip_url=module_report_zip_url,
            module_report_folder_url=module_report_folder_url,
        ),
        _base_test_sheet_spec(task, patches),
    ]


def _ensure_task_feishu_folder(task: Dict[str, Any]) -> Dict[str, str]:
    """Create or reuse ``{configuredRoot}/{fromTag-toTag_执行时间}/`` for this task."""
    from web_app.feishu_service import ensure_task_result_folder, task_run_stamp

    return ensure_task_result_folder(
        str(task.get("from_tag") or ""),
        str(task.get("to_tag") or ""),
        existing_token=str(task.get("feishu_task_folder_token") or ""),
        run_stamp=task_run_stamp(task),
    )


def _attach_module_diff_report(
    task_id: str,
    *,
    folder_token: str = "",
    folder_url: str = "",
) -> Dict[str, Any]:
    """Build + upload offline module HTML report; best-effort (never blocks sheet export)."""
    out: Dict[str, Any] = {
        "module_report_urls": {},
        "subtag_report_urls": {},
        "module_report_zip_url": "",
        "module_report_folder_url": folder_url or "",
        "module_report_zip_path": "",
        "folder_token": folder_token or "",
    }
    try:
        from web_app.module_diff_report import (
            build_module_diff_report,
            module_url_map,
            subtag_url_map,
            upload_module_diff_report,
        )
        report = build_module_diff_report(task_id, write_zip=False)
        out["module_report_zip_path"] = str(report.report_dir)
        print(f"[REPORT] 模块代码详情已生成: {report.report_dir}", flush=True)
        try:
            uploaded = upload_module_diff_report(
                report,
                folder_token=folder_token,
                folder_url=folder_url,
            )
            out["module_report_urls"] = module_url_map(uploaded)
            out["subtag_report_urls"] = subtag_url_map(uploaded)
            out["module_report_zip_url"] = ""
            out["module_report_folder_url"] = uploaded.folder_url or folder_url or ""
            out["folder_token"] = uploaded.folder_token or folder_token or ""
            print(
                f"[FEISHU] 模块代码详情已上传: folder={out['module_report_folder_url']} "
                f"modules={len(out['module_report_urls'])} "
                f"subtags={len(out['subtag_report_urls'])}",
                flush=True,
            )
        except Exception as exc:
            print(f"[FEISHU] 模块代码详情上传失败（本地报告仍可用）: {exc}", flush=True)
    except Exception as exc:
        print(f"[REPORT] 模块代码详情生成失败: {exc}", flush=True)
    return out


def export_checklist_to_feishu(task_id: str) -> Dict[str, str]:
    """Create a new Feishu spreadsheet for this task and write checklist data."""
    from web_app.database import update_task
    from web_app.feishu_service import (
        FeishuError,
        add_sheet,
        create_spreadsheet,
        list_sheets,
        transfer_spreadsheet_owner,
        update_sheet_props,
    )

    task = get_task(task_id)
    if not task:
        raise FileNotFoundError("任务不存在")

    folder = _ensure_task_feishu_folder(task)
    folder_token = str(folder.get("token") or "")
    folder_url = str(folder.get("url") or "")

    report_info = _attach_module_diff_report(
        task_id,
        folder_token=folder_token,
        folder_url=folder_url,
    )
    title, sheets = build_checklist_sheets(
        task_id,
        module_report_urls=report_info.get("module_report_urls") or {},
        subtag_report_urls=report_info.get("subtag_report_urls") or {},
        module_report_zip_url=str(report_info.get("module_report_zip_url") or ""),
        module_report_folder_url=folder_url or str(report_info.get("module_report_folder_url") or ""),
    )
    created = create_spreadsheet(title, folder_token=folder_token)
    token = created["token"]
    url = created["url"]

    existing = list_sheets(token)
    first_id = existing[0]["sheet_id"] if existing else ""
    if not first_id:
        raise FeishuError("新建表格后未能读取工作表 ID")

    overview, base = sheets
    update_sheet_props(
        token,
        first_id,
        title=overview.title,
        frozen_row_count=overview.frozen_rows,
    )
    _write_sheet_to_feishu(token, first_id, overview)

    second_id = add_sheet(token, base.title, index=1)
    if not second_id:
        raise FeishuError("创建「基础测试需求」工作表失败")
    _write_sheet_to_feishu(token, second_id, base)

    try:
        transfer_spreadsheet_owner(token)
    except Exception as exc:
        print(f"[FEISHU] 转移所有者失败: {exc}", flush=True)

    update_task(
        task_id,
        feishu_checklist_url=url,
        feishu_checklist_token=token,
        feishu_module_report_url=str(report_info.get("module_report_folder_url") or folder_url or ""),
        feishu_module_report_folder=folder_url or str(report_info.get("module_report_folder_url") or ""),
        feishu_task_folder_token=folder_token,
    )
    return {
        "url": url,
        "token": token,
        "title": title,
        "module_report_url": str(report_info.get("module_report_folder_url") or folder_url or ""),
        "module_report_folder": folder_url or str(report_info.get("module_report_folder_url") or ""),
        "module_report_zip_path": str(report_info.get("module_report_zip_path") or ""),
        "task_folder_token": folder_token,
    }


def try_export_checklist_to_feishu(task_id: str) -> Optional[Dict[str, str]]:
    """Best-effort Feishu export used after analysis completes."""
    try:
        result = export_checklist_to_feishu(task_id)
        print(f"[FEISHU] 验收表格已创建: {result.get('url')}", flush=True)
        return result
    except Exception as exc:
        print(f"[FEISHU] 导出飞书表格失败: {exc}", flush=True)
        return None


def _overview_sheet_spec(
    task: Dict[str, Any],
    patches: List[Dict[str, Any]],
    *,
    module_report_urls: Optional[Dict[str, str]] = None,
    subtag_report_urls: Optional[Dict[str, str]] = None,
    module_report_zip_url: str = "",
    module_report_folder_url: str = "",
) -> SheetSpec:
    from web_app.feishu_service import feishu_link_cell

    from_tag = str(task.get("from_tag") or "").strip()
    to_tag = str(task.get("to_tag") or "").strip()
    last_col = len(OVERVIEW_HEADERS)
    report_urls = module_report_urls or {}
    subtag_urls = subtag_report_urls or {}
    report_link = module_report_folder_url or module_report_zip_url or str(task.get("feishu_module_report_folder") or task.get("feishu_module_report_url") or "")
    if report_link:
        report_cell: Any = feishu_link_cell("打开模块代码详情（云盘 HTML）", report_link)
    else:
        report_cell = "导出飞书时将自动生成模块代码详情报告"
    meta_rows = [
        ("GKI升级版本", f"{from_tag} -> {to_tag}" if from_tag and to_tag else to_tag or from_tag),
        ("代码审核原则", _HEADER_PRINCIPLES["代码审核原则"]),
        ("测试与验证策略", _compose_test_strategy(task)),
        ("质量控制", _HEADER_PRINCIPLES["质量控制"]),
        ("质量定责", _HEADER_PRINCIPLES["质量定责"]),
        ("转测版本路径", ""),
        ("模块代码详情", report_cell),
    ]
    values: List[List[Any]] = []
    merges: List[Tuple[int, int, int, int]] = []
    for idx, (label, value) in enumerate(meta_rows, start=1):
        values.append(_pad_row([label, value], last_col))
        merges.append((idx, 2, idx, last_col))
    values.append(list(OVERVIEW_HEADERS))
    header_row = len(values)

    owners = load_module_owners()
    blocks = _build_overview_blocks(task, patches, owners)
    skip_rows: List[int] = []
    tag_blocks: List[Tuple[int, int, str]] = []
    cursor = header_row + 1
    for block in blocks:
        start = cursor
        kind = str(block.get("kind") or "")
        subtag_key = str(block.get("subtag_key") or "").strip()
        subtag_url = subtag_urls.get(subtag_key) if kind == "subtag" else ""
        for row_i, row in enumerate(block["rows"]):
            cells = _pad_row(row.to_values(), last_col)
            link_idx = _COL_MODULE_LINK - 1
            if kind == "subtag":
                # 子 tag 整段共用一个详情链接（列合并）
                if row_i == 0 and link_idx < len(cells):
                    if subtag_url:
                        cells[link_idx] = feishu_link_cell("打开子tag详情", subtag_url)
                    else:
                        cells[link_idx] = "—"
                elif link_idx < len(cells):
                    cells[link_idx] = ""
            else:
                mod_url = report_urls.get(row.module_name) or ""
                if mod_url and not row.skipped and link_idx < len(cells):
                    cells[link_idx] = feishu_link_cell("打开详情", mod_url)
                elif not row.skipped and link_idx < len(cells):
                    cells[link_idx] = "—"
            values.append(_wrap_row(cells, _OVERVIEW_WRAP))
            if row.skipped:
                skip_rows.append(cursor)
            cursor += 1
        end = cursor - 1
        if end >= start:
            merges.append((start, _COL_TAG, end, _COL_TAG_SPAN))
            if block.get("merge_submit"):
                merges.append((start, _COL_SUBMIT, end, _COL_SUBMIT))
            if kind == "subtag":
                merges.append((start, _COL_MODULE_LINK, end, _COL_MODULE_LINK))
            merges.extend(
                (span_start, _COL_DEPT, span_end, _COL_DEPT)
                for span_start, span_end in _department_spans(block["rows"], start)
            )
            tag_blocks.append((start, end, kind))
    return SheetSpec(
        title="版本验收总览",
        values=values,
        merges=_unique_merges(merges),
        frozen_rows=header_row,
        col_widths=_COL_WIDTHS_OVERVIEW,
        col_pixels=_COL_PIXELS_OVERVIEW,
        header_row=header_row,
        skip_rows=skip_rows,
        tag_blocks=tag_blocks,
        center_cols=list(_OVERVIEW_CENTER_COLS),
        row_heights=_overview_row_heights(values, header_row, skip_rows),
    )


def _base_test_sheet_spec(task: Dict[str, Any], patches: List[Dict[str, Any]]) -> SheetSpec:
    owners = load_module_owners()
    modules_doc = _load_json_field(task.get("module_summaries"))
    modules_map = (modules_doc or {}).get("modules") or {}
    grouped = _group_patches(patches)
    last_col = len(BASE_TEST_HEADERS)
    values: List[List[str]] = [list(BASE_TEST_HEADERS)]
    skip_rows: List[int] = []
    names = list(get_dev_module_names())
    extra = [name for name in grouped if name not in names]
    row_idx = 2
    last_dept = ""
    dept_start = 2
    merges: List[Tuple[int, int, int, int]] = []
    for name in names + extra:
        owner = owners.get(name, ModuleOwner())
        analysis = ((modules_map.get(name) or {}).get("analysis") or {}) if isinstance(modules_map.get(name), dict) else {}
        mod_patches = grouped.get(name) or []
        skipped = not mod_patches
        tests = _export_test_suggestions(task, name, analysis, mod_patches) if not skipped else ""
        values.append(_wrap_row(_pad_row([
            _department_label(owner),
            name,
            owner.owner,
            "否" if skipped else "是",
            tests,
            "",
            "",
        ], last_col), _BASE_WRAP))
        if skipped:
            skip_rows.append(row_idx)
        dept = owner.department or ""
        if dept != last_dept and last_dept and row_idx - 1 > dept_start:
            merges.append((dept_start, 1, row_idx - 1, 1))
        if dept != last_dept:
            dept_start = row_idx
            last_dept = dept
        row_idx += 1
    if last_dept and row_idx - 1 > dept_start:
        merges.append((dept_start, 1, row_idx - 1, 1))
    return SheetSpec(
        title="基础测试需求",
        values=values,
        merges=_unique_merges(merges),
        frozen_rows=1,
        col_widths=_COL_WIDTHS_BASE,
        col_pixels=_COL_PIXELS_BASE,
        header_row=1,
        skip_rows=skip_rows,
        center_cols=list(_BASE_CENTER_COLS),
        row_heights=_base_row_heights(values, skip_rows),
    )


def refresh_feishu_checklist(task_id: str) -> Dict[str, str]:
    """Rewrite checklist content/styles into the task's existing Feishu spreadsheet."""
    from web_app.database import update_task
    from web_app.feishu_service import FeishuError, list_sheets

    task = get_task(task_id)
    if not task:
        raise FileNotFoundError("任务不存在")
    token = str(task.get("feishu_checklist_token") or "").strip()
    url = str(task.get("feishu_checklist_url") or "").strip()
    if not token:
        return export_checklist_to_feishu(task_id)

    folder = _ensure_task_feishu_folder(task)
    folder_token = str(folder.get("token") or "")
    folder_url = str(folder.get("url") or "")

    report_info = _attach_module_diff_report(
        task_id,
        folder_token=folder_token,
        folder_url=folder_url,
    )
    title, sheets = build_checklist_sheets(
        task_id,
        module_report_urls=report_info.get("module_report_urls") or {},
        subtag_report_urls=report_info.get("subtag_report_urls") or {},
        module_report_zip_url=str(report_info.get("module_report_zip_url") or ""),
        module_report_folder_url=folder_url or str(report_info.get("module_report_folder_url") or ""),
    )
    existing = list_sheets(token)
    if not existing:
        raise FeishuError("已有飞书表格中未找到工作表")
    by_title = {item["title"]: item["sheet_id"] for item in existing}
    for idx, spec in enumerate(sheets):
        sheet_id = by_title.get(spec.title) or (existing[idx]["sheet_id"] if idx < len(existing) else "")
        if not sheet_id:
            raise FeishuError(f"未找到工作表：{spec.title}")
        _write_sheet_to_feishu(token, sheet_id, spec, merge=True)
    if not url:
        url = f"https://transsioner.feishu.cn/sheets/{token}"
        update_task(task_id, feishu_checklist_url=url)
    update_task(
        task_id,
        feishu_module_report_url=str(report_info.get("module_report_folder_url") or folder_url or ""),
        feishu_module_report_folder=folder_url or str(report_info.get("module_report_folder_url") or ""),
        feishu_task_folder_token=folder_token,
    )
    return {
        "url": url,
        "token": token,
        "title": title,
        "module_report_url": str(report_info.get("module_report_folder_url") or folder_url or ""),
        "module_report_folder": folder_url or str(report_info.get("module_report_folder_url") or ""),
        "module_report_zip_path": str(report_info.get("module_report_zip_path") or ""),
        "task_folder_token": folder_token,
    }


def _write_sheet_to_feishu(token: str, sheet_id: str, spec: SheetSpec, *, merge: bool = True) -> None:
    from web_app.feishu_service import add_columns, add_rows, merge_cells, unmerge_cells, write_values

    nrows = len(spec.values)
    ncols = max((len(row) for row in spec.values), default=1)
    min_cols = max(12, ncols)
    if min_cols > 20:
        add_columns(token, sheet_id, min_cols - 20)
    if nrows > 200:
        add_rows(token, sheet_id, min(5000, nrows - 200 + 20))
    a1 = f"{sheet_id}!A1:{get_column_letter(ncols)}{max(nrows, 1)}"
    write_values(token, a1, spec.values)
    try:
        unmerge_cells(token, f"{sheet_id}!A1:{get_column_letter(max(ncols, 12))}{max(nrows, 1)}")
    except Exception as exc:
        print(f"[FEISHU] 取消合并失败: {exc}", flush=True)
    if merge:
        for start_r, start_c, end_r, end_c in spec.merges:
            if end_r <= start_r and end_c <= start_c:
                continue
            rng = f"{sheet_id}!{get_column_letter(start_c)}{start_r}:{get_column_letter(end_c)}{end_r}"
            try:
                merge_cells(token, rng)
            except Exception as exc:
                print(f"[FEISHU] 合并单元格失败 {rng}: {exc}", flush=True)
    _apply_feishu_styles(token, sheet_id, spec, ncols, nrows)
    _apply_feishu_layout(token, sheet_id, spec, ncols, nrows)
    _apply_feishu_dropdowns(token, sheet_id, spec, nrows)


def _feishu_color(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return text if text.startswith("#") else f"#{text}"


def _feishu_v2_style(
    *,
    back_color: str = "",
    bold: bool = False,
    fore_color: str = "",
    h_align: str = "LEFT",
    v_align: str = "TOP",
    font_size: str = "10pt/1.5",
    border: bool = False,
) -> Dict[str, Any]:
    align_h = {"LEFT": 0, "CENTER": 1, "RIGHT": 2}.get(h_align.upper(), 0)
    align_v = {"TOP": 0, "CENTER": 1, "MIDDLE": 1, "BOTTOM": 2}.get(v_align.upper(), 0)
    style: Dict[str, Any] = {
        "font": {"bold": bold, "fontSize": font_size},
        "hAlign": align_h,
        "vAlign": align_v,
    }
    if back_color:
        style["backColor"] = _feishu_color(back_color)
    if fore_color:
        style["foreColor"] = _feishu_color(fore_color)
    if border:
        style["borderType"] = "FULL_BORDER"
        style["borderColor"] = _COLOR_BORDER
    return style


def _apply_feishu_styles(token: str, sheet_id: str, spec: SheetSpec, ncols: int, nrows: int) -> None:
    from web_app.feishu_service import styles_batch_update

    last = get_column_letter(ncols)
    items: List[Dict[str, Any]] = []
    if nrows >= 1:
        items.append({
            "ranges": [f"{sheet_id}!A1:{last}{nrows}"],
            "style": _feishu_v2_style(h_align="LEFT", v_align="TOP", border=True),
        })
    if spec.header_row > 1 and spec.title == "版本验收总览":
        # 元信息：标签列居中加粗，内容左对齐垂直居中（对齐模板）
        items.append({
            "ranges": [f"{sheet_id}!A1:A{spec.header_row - 1}"],
            "style": _feishu_v2_style(bold=True, h_align="CENTER", v_align="CENTER", font_size="10pt/1.5"),
        })
        items.append({
            "ranges": [f"{sheet_id}!B1:{last}{spec.header_row - 1}"],
            "style": _feishu_v2_style(h_align="LEFT", v_align="CENTER", font_size="10pt/1.5"),
        })
    header_fg = _COLOR_HEADER_FG
    header_bg = _COLOR_HEADER_BG
    if spec.title == "基础测试需求":
        header_fg = _COLOR_BASE_HEADER_FG
        header_bg = ""
    items.append({
        "ranges": [f"{sheet_id}!A{spec.header_row}:{last}{spec.header_row}"],
        "style": _feishu_v2_style(
            back_color=header_bg,
            bold=spec.title == "版本验收总览",
            fore_color=header_fg,
            h_align="CENTER",
            v_align="CENTER",
            font_size="10pt/1.5",
            border=True,
        ),
    })
    if spec.title == "基础测试需求" and spec.header_row >= 1:
        letter = get_column_letter(len(BASE_TEST_HEADERS))
        row = spec.header_row
        items.append({
            "ranges": [f"{sheet_id}!{letter}{row}:{letter}{row}"],
            "style": _feishu_v2_style(
                back_color=_COLOR_BASE_HEADER_ACCENT,
                fore_color=_COLOR_HEADER_FG,
                h_align="LEFT",
                v_align="CENTER",
                font_size="11pt/1.5",
                border=True,
            ),
        })
    for col in spec.center_cols:
        letter = get_column_letter(col)
        start = spec.header_row + 1
        if start <= nrows:
            items.append({
                "ranges": [f"{sheet_id}!{letter}{start}:{letter}{nrows}"],
                "style": _feishu_v2_style(h_align="CENTER", v_align="CENTER"),
            })
    try:
        styles_batch_update(token, items)
    except Exception as exc:
        print(f"[FEISHU] 批量样式失败: {exc}", flush=True)


def _apply_feishu_dropdowns(token: str, sheet_id: str, spec: SheetSpec, nrows: int) -> None:
    if spec.title != "版本验收总览":
        return
    start_row = spec.header_row + 1
    if start_row > nrows:
        return
    from web_app.feishu_service import set_dropdown

    def col_range(col: int) -> str:
        letter = get_column_letter(col)
        return f"{sheet_id}!{letter}{start_row}:{letter}{nrows}"

    try:
        set_dropdown(
            token, col_range(_COL_REVIEW_PLUS), _REVIEW_PLUS_OPTIONS,
            colors=_REVIEW_PLUS_COLORS, highlight=True,
        )
        set_dropdown(
            token, col_range(_COL_NEGATIVE), _YES_NO_NA_OPTIONS,
            colors=_YES_NO_NA_COLORS, highlight=True,
        )
        set_dropdown(
            token, col_range(_COL_MUST_MERGE), _YES_NO_NA_OPTIONS,
            colors=_YES_NO_NA_COLORS, highlight=True,
        )
    except Exception as exc:
        print(f"[FEISHU] 下拉框设置失败: {exc}", flush=True)


def _apply_xlsx_dropdowns(ws: Worksheet, header_row: int, last_row: int) -> None:
    if last_row <= header_row:
        return

    def add_list(col: int, options: Sequence[str]) -> None:
        letter = get_column_letter(col)
        formula = '"' + ",".join(options) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(f"{letter}{header_row + 1}:{letter}{last_row}")
        ws.add_data_validation(dv)

    add_list(_COL_REVIEW_PLUS, _REVIEW_PLUS_OPTIONS)
    add_list(_COL_NEGATIVE, _YES_NO_NA_OPTIONS)
    add_list(_COL_MUST_MERGE, _YES_NO_NA_OPTIONS)


def _apply_feishu_layout(token: str, sheet_id: str, spec: SheetSpec, ncols: int, nrows: int) -> None:
    from web_app.feishu_service import update_dimension, update_sheet_props

    pixels = spec.col_pixels or {
        col: int(width * 7.2 + 8) for col, width in spec.col_widths.items()
    }
    for col, size in pixels.items():
        _update_dimension(token, sheet_id, "COLUMNS", col, col, size)
    for start, end, size in _iter_index_groups(spec.row_heights):
        _update_dimension(token, sheet_id, "ROWS", start, end, size)
    try:
        visible_cols = max(ncols, len(OVERVIEW_HEADERS))
        update_dimension(token, sheet_id, "COLUMNS", 1, visible_cols, visible=True)
    except Exception as exc:
        print(f"[FEISHU] 显示数据列失败: {exc}", flush=True)
    if ncols < 20:
        try:
            update_dimension(
                token, sheet_id, "COLUMNS", max(ncols, len(OVERVIEW_HEADERS)) + 1, 20, visible=False,
            )
        except Exception as exc:
            print(f"[FEISHU] 隐藏多余列失败: {exc}", flush=True)
    try:
        update_sheet_props(token, sheet_id, frozen_row_count=spec.frozen_rows or 0)
    except Exception as exc:
        print(f"[FEISHU] 更新冻结行失败: {exc}", flush=True)


def _update_dimension(token: str, sheet_id: str, dimension: str, start: int, end: int, size: int) -> None:
    from web_app.feishu_service import update_dimension

    try:
        update_dimension(token, sheet_id, dimension, start, end, fixed_size=int(size))
    except Exception as exc:
        print(f"[FEISHU] 行列尺寸失败 {dimension} {start}:{end}: {exc}", flush=True)


def _iter_index_groups(mapping: Dict[int, int]) -> List[Tuple[int, int, int]]:
    if not mapping:
        return []
    items = sorted(mapping.items())
    groups: List[Tuple[int, int, int]] = []
    start, prev, size = items[0][0], items[0][0], items[0][1]
    for idx, val in items[1:]:
        if idx == prev + 1 and val == size:
            prev = idx
            continue
        groups.append((start, prev, size))
        start, prev, size = idx, idx, val
    groups.append((start, prev, size))
    return groups


def _sheet_list(info: Dict[str, Any]) -> List[Dict[str, str]]:
    sheets = info.get("sheets") or info.get("sheet") or []
    if isinstance(sheets, dict):
        sheets = sheets.get("sheets") or []
    out: List[Dict[str, str]] = []
    if not isinstance(sheets, list):
        return out
    for item in sheets:
        if not isinstance(item, dict):
            continue
        sheet_id = str(item.get("sheet_id") or item.get("sheetId") or "").strip()
        title = str(item.get("title") or "").strip()
        if sheet_id:
            out.append({"sheet_id": sheet_id, "title": title})
    return out


def _extract_sheet_id(payload: Dict[str, Any]) -> str:
    for key in ("sheet_id", "sheetId"):
        value = payload.get(key)
        if value:
            return str(value)
    sheet = payload.get("sheet")
    if isinstance(sheet, dict):
        return str(sheet.get("sheet_id") or sheet.get("sheetId") or "")
    return ""


def _sheet_id_by_title(info: Dict[str, Any], title: str) -> str:
    for item in _sheet_list(info):
        if item.get("title") == title:
            return item["sheet_id"]
    return ""


def _cell_as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return str(value.get("text") or value.get("link") or "")
    return str(value)


def _pad_row(values: Sequence[Any], width: int) -> List[Any]:
    row: List[Any] = []
    for v in values:
        if v is None:
            row.append("")
        elif isinstance(v, dict):
            row.append(v)
        else:
            row.append(v)
    if len(row) < width:
        row.extend([""] * (width - len(row)))
    return row[:width]


def _department_spans(rows: List[ChecklistRow], start: int) -> List[Tuple[int, int]]:
    if not rows:
        return []
    spans: List[Tuple[int, int]] = []
    run_dept = rows[0].department
    run_start = start
    for offset, row in enumerate(rows[1:], start=1):
        if row.department == run_dept:
            continue
        end = start + offset - 1
        if end > run_start:
            spans.append((run_start, end))
        run_dept = row.department
        run_start = start + offset
    end = start + len(rows) - 1
    if end > run_start:
        spans.append((run_start, end))
    return spans


def _unique_merges(merges: List[Tuple[int, int, int, int]]) -> List[Tuple[int, int, int, int]]:
    out: List[Tuple[int, int, int, int]] = []
    seen = set()
    for item in merges:
        if item in seen:
            continue
        start_r, start_c, end_r, end_c = item
        if end_r < start_r or end_c < start_c:
            continue
        if end_r == start_r and end_c == start_c:
            continue
        seen.add(item)
        out.append(item)
    return out


def _load_checklist_config() -> Dict[str, Any]:
    """Load unified checklist config from config/checklist_config.yaml."""
    path = checklist_config_path()
    if not path.is_file():
        return {}
    try:
        raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except Exception:
        return {}
    return raw if isinstance(raw, dict) else {}


def load_integration_maintenance() -> Dict[str, Any]:
    """Load fixed 集成维护 checklist rows from checklist_config.yaml."""
    block = _load_checklist_config().get("integration_maintenance")
    return block if isinstance(block, dict) else {}


def _integration_maintenance_rows() -> List[ChecklistRow]:
    """Fixed integration-maintenance checks after overall module rows (see Feishu template)."""
    cfg = load_integration_maintenance()
    if not cfg or not cfg.get("enabled", True):
        return []
    dept = str(cfg.get("department") or "集成维护").strip()
    leads = str(cfg.get("leads") or "").strip()
    department = f"{dept}@{leads}" if leads else dept

    rows: List[ChecklistRow] = []
    modules = cfg.get("modules") or {}
    if not isinstance(modules, dict):
        return []
    for module_name, info in modules.items():
        if not isinstance(info, dict):
            info = {}
        name = str(module_name or "").strip()
        owner = str(info.get("owner") or "").strip()
        if not name:
            continue
        domain = f"{name}-{owner}" if owner else name
        rows.append(
            ChecklistRow(
                module_name=name,
                department=department,
                domain_owner=domain,
                reviewer=str(info.get("reviewer") or "").strip() or _integration_reviewer_from_domain(domain),
                review_plus=str(info.get("review_plus") or "").strip(),
                modification=str(info.get("modification") or "").strip(),
                focus_review=str(info.get("focus_review") or "").strip(),
                negative_impact=str(info.get("negative_impact") or "").strip(),
                risk_description=str(info.get("risk_description") or "").strip(),
                must_merge=str(info.get("must_merge") or "").strip(),
                test_suggestions=str(info.get("test_suggestions") or "").strip(),
                block_kind="integration",
                skipped=False,
            )
        )
    return rows


def _integration_reviewer_from_domain(domain: str) -> str:
    text = str(domain or "").strip()
    if not text:
        return ""
    if "-" in text:
        text = text.split("-", 1)[1].strip()
    for mark in ("（", "("):
        if mark in text:
            text = text.split(mark, 1)[0].strip()
    return text


def load_module_owners() -> Dict[str, ModuleOwner]:
    mapping: Dict[str, ModuleOwner] = {}
    departments = _load_checklist_config().get("departments") or {}
    if not isinstance(departments, dict):
        return mapping
    for dept_name, dept in departments.items():
        if not isinstance(dept, dict):
            continue
        lead = str(dept.get("lead") or "").strip()
        modules = dept.get("modules") or {}
        if not isinstance(modules, dict):
            continue
        for mod_name, info in modules.items():
            info = info or {}
            if not isinstance(info, dict):
                info = {}
            mapping[str(mod_name)] = ModuleOwner(
                department=str(dept_name),
                lead=lead,
                owner=str(info.get("owner") or "").strip(),
                reviewer=str(info.get("reviewer") or "").strip(),
            )
    return mapping


def _hydrate_patch(row: Dict[str, Any]) -> Dict[str, Any]:
    data_raw = row.get("data")
    full: Dict[str, Any] = {}
    if data_raw:
        try:
            loaded = json.loads(data_raw) if isinstance(data_raw, str) else dict(data_raw)
            if isinstance(loaded, dict):
                full = loaded
        except (json.JSONDecodeError, TypeError):
            full = {}
    merged = dict(full)
    for key in ("sha", "title", "score", "recommendation", "key_risk", "dev_module"):
        value = row.get(key)
        if value not in (None, ""):
            merged[key] = value
    if not merged.get("sha"):
        merged["sha"] = row.get("sha") or full.get("commit") or full.get("sha") or ""
    if not merged.get("commit"):
        merged["commit"] = merged.get("sha")
    merged["dev_module"] = str(merged.get("dev_module") or "其他").strip() or "其他"
    return merged


def _write_overview_sheet(
    ws: Worksheet,
    task: Dict[str, Any],
    patches: List[Dict[str, Any]],
) -> None:
    from_tag = str(task.get("from_tag") or "").strip()
    to_tag = str(task.get("to_tag") or "").strip()
    repo = str(task.get("repo") or "").rstrip("/")

    meta_rows = [
        ("GKI升级版本", f"{from_tag} -> {to_tag}" if from_tag and to_tag else to_tag or from_tag),
        ("代码审核原则", _HEADER_PRINCIPLES["代码审核原则"]),
        ("测试与验证策略", _compose_test_strategy(task)),
        ("质量控制", _HEADER_PRINCIPLES["质量控制"]),
        ("质量定责", _HEADER_PRINCIPLES["质量定责"]),
        ("转测版本路径", ""),
        ("模块代码详情", "表格仅列关键摘要与重点 Review；点「模块详情」打开 HTML 看完整解读与 diff"),
    ]
    last_col = len(OVERVIEW_HEADERS)
    for idx, (label, value) in enumerate(meta_rows, start=1):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        if last_col > 2:
            ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=last_col)
        for col in range(1, last_col + 1):
            cell = ws.cell(idx, col)
            if col == 1:
                cell.font = _FONT_LABEL
                cell.alignment = _ALIGN_META_LABEL
            else:
                cell.font = _FONT_BODY
                cell.alignment = _ALIGN_META_VALUE
            cell.border = _THIN
        ws.row_dimensions[idx].height = 22 if idx == 1 else 28

    header_row = len(meta_rows) + 1
    # 表头不合并 A:B（与模板一致：A 空、B 写 Tag 标题），数据区再合并 Tag
    for col, title in enumerate(OVERVIEW_HEADERS, start=1):
        cell = ws.cell(header_row, col, title)
        cell.font = _FONT_HEADER
        cell.fill = _FILL_HEADER
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = f"A{header_row + 1}"

    owners = load_module_owners()
    blocks = _build_overview_blocks(task, patches, owners)
    cursor = header_row + 1
    for block in blocks:
        start = cursor
        for row in block["rows"]:
            values = row.to_values()
            for col, value in enumerate(values, start=1):
                cell = ws.cell(cursor, col, value)
                cell.font = _FONT_BODY
                cell.alignment = _ALIGN_CENTER if col in _OVERVIEW_CENTER_COLS else _ALIGN_WRAP
                cell.border = _THIN
            ws.row_dimensions[cursor].height = _row_height(values)
            cursor += 1
        end = cursor - 1
        if end >= start:
            ws.merge_cells(
                start_row=start, start_column=_COL_TAG,
                end_row=end, end_column=_COL_TAG_SPAN,
            )
            if block.get("merge_submit"):
                _merge_column(ws, _COL_SUBMIT, start, end)
            _merge_same_department(ws, block["rows"], start, col=_COL_DEPT)

    for col, width in _COL_WIDTHS_OVERVIEW.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"1:{header_row}"
    _apply_xlsx_dropdowns(ws, header_row, cursor - 1)


def _write_base_test_sheet(
    ws: Worksheet,
    task: Dict[str, Any],
    patches: List[Dict[str, Any]],
) -> None:
    owners = load_module_owners()
    modules_doc = _load_json_field(task.get("module_summaries"))
    modules_map = (modules_doc or {}).get("modules") or {}
    grouped = _group_patches(patches)

    for col, title in enumerate(BASE_TEST_HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.font = _FONT_BASE_HEADER
        if col == len(BASE_TEST_HEADERS):
            cell.fill = _FILL_BASE_HEADER_ACCENT
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"

    names = list(get_dev_module_names())
    extra = [name for name in grouped if name not in names]
    row_idx = 2
    last_dept = ""
    dept_start = 2
    for name in names + extra:
        owner = owners.get(name, ModuleOwner())
        analysis = ((modules_map.get(name) or {}).get("analysis") or {}) if isinstance(modules_map.get(name), dict) else {}
        mod_patches = grouped.get(name) or []
        skipped = not mod_patches
        tests = _export_test_suggestions(task, name, analysis, mod_patches)
        must = "否" if skipped else "是"
        values = [
            _department_label(owner),
            name,
            owner.owner,
            must,
            "" if skipped else tests,
            "",
            "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col, value)
            cell.font = _FONT_BODY
            cell.alignment = _ALIGN_WRAP
            cell.border = _THIN
        ws.row_dimensions[row_idx].height = _row_height(values, min_h=22, max_h=90)
        dept = owner.department or ""
        if dept != last_dept and last_dept:
            if row_idx - 1 > dept_start:
                _merge_column(ws, 1, dept_start, row_idx - 1)
        if dept != last_dept:
            dept_start = row_idx
            last_dept = dept
        row_idx += 1
    if last_dept and row_idx - 1 > dept_start:
        _merge_column(ws, 1, dept_start, row_idx - 1)

    for col, width in _COL_WIDTHS_BASE.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _build_overview_blocks(
    task: Dict[str, Any],
    patches: List[Dict[str, Any]],
    owners: Dict[str, ModuleOwner],
) -> List[Dict[str, Any]]:
    from_tag = str(task.get("from_tag") or "").strip()
    to_tag = str(task.get("to_tag") or "").strip()
    repo = str(task.get("repo") or "").rstrip("/")
    modules_doc = _load_json_field(task.get("module_summaries"))
    modules_map = (modules_doc or {}).get("modules") or {}
    grouped = _group_patches(patches)

    overall_rows = _overall_module_rows(
        from_tag, to_tag, repo, task, grouped, modules_map, owners,
    )
    integration_rows = _integration_maintenance_rows()
    if integration_rows:
        overall_rows.extend(integration_rows)
    blocks: List[Dict[str, Any]] = [{
        "kind": "overall",
        "merge_submit": True,
        "rows": overall_rows,
    }]
    for sub in _iter_incremental_subtags(task, patches):
        rows = _subtag_module_rows(sub, owners)
        if rows:
            blocks.append({
                "kind": "subtag",
                "merge_submit": False,
                "subtag_key": str(sub.get("to_tag") or "").strip(),
                "rows": rows,
            })
    _polish_checklist_blocks(task, blocks, patches)
    return blocks


def _polish_checklist_blocks(
    task: Dict[str, Any],
    blocks: List[Dict[str, Any]],
    patches: Optional[List[Dict[str, Any]]] = None,
) -> None:
    rows: List[ChecklistRow] = []
    for block in blocks:
        for row in block.get("rows") or []:
            if isinstance(row, ChecklistRow) and not row.skipped:
                if (row.block_kind or "") == "integration":
                    continue
                rows.append(row)
    if not rows:
        return

    grouped = _group_patches(patches or [])
    modules_doc = _load_json_field(task.get("module_summaries"))
    modules_map = (modules_doc or {}).get("modules") or {}

    items = []
    for idx, row in enumerate(rows):
        scope = (row.tag_text or "").split("\n")[0].strip() or "整体升级"
        is_subtag = (row.block_kind or "") == "subtag"
        # 子 tag：source_patches 仅为本段相邻增量中该模块的 patch，禁止回落到总包模块分析
        if row.source_patches:
            mod_patches = list(row.source_patches)
        elif is_subtag:
            mod_patches = []
        else:
            mod_patches = grouped.get(row.module_name) or []
        analysis: Dict[str, Any] = {}
        entry = modules_map.get(row.module_name)
        if (not is_subtag) and isinstance(entry, dict):
            analysis = entry.get("analysis") or {}
        stats = _module_stats_line(mod_patches)
        file_locations = _module_file_locations_summary(
            mod_patches,
            code_tree=None if is_subtag else (
                entry.get("code_tree") if isinstance(entry, dict) else None
            ),
            limit_dirs=4,
            limit_files=0,
        )
        cumulative_scope = ""
        if is_subtag:
            for line in str(row.tag_text or "").splitlines():
                text = line.strip()
                if text.startswith("Diff:") or text.startswith("结论:"):
                    cumulative_scope = text
                    break
            if not cumulative_scope:
                cumulative_scope = scope
        items.append({
            "id": str(idx),
            "module": row.module_name,
            "scope": scope,
            "block_kind": "subtag" if is_subtag else "overall",
            "cumulative_scope": cumulative_scope,
            "patch_count": len(mod_patches),
            "stats": stats,
            "file_locations": file_locations,
            "analysis_context": "" if is_subtag else _analysis_context_for_polish(analysis),
            "patch_digest": _patch_digest_for_checklist(mod_patches, limit=20 if is_subtag else 12),
            "modification": row.modification,
            "focus_review": row.focus_review,
            "test_suggestions": row.test_suggestions,
            "negative_impact": row.negative_impact,
            "risk_description": row.risk_description,
            "must_merge": row.must_merge,
        })
    polished: Dict[str, Dict[str, str]] = {}
    try:
        from web_app.llm_service import is_llm_configured, polish_checklist_cells
        if is_llm_configured():
            overall_items = [i for i in items if i.get("block_kind") != "subtag"]
            subtag_items = [i for i in items if i.get("block_kind") == "subtag"]
            print(
                f"[CHECKLIST] 正在用大模型打磨表格文案："
                f"总包 {len(overall_items)} + 子tag {len(subtag_items)}（分批，防复用）…",
                flush=True,
            )
            if overall_items:
                polished.update(
                    polish_checklist_cells(task, overall_items, batch_label="总包") or {},
                )
            if subtag_items:
                polished.update(
                    polish_checklist_cells(task, subtag_items, batch_label="子tag") or {},
                )
        else:
            print("[CHECKLIST] 未配置 LLM，表格文案仅做压缩", flush=True)
    except Exception as exc:
        print(f"[CHECKLIST] 表格打磨跳过: {exc}", flush=True)

    overall_polished_by_module: Dict[str, Dict[str, str]] = {}
    for idx, row in enumerate(rows):
        if (row.block_kind or "") == "subtag":
            continue
        got = polished.get(str(idx)) or {}
        overall_polished_by_module[row.module_name] = {
            "modification": got.get("modification") or row.modification,
            "focus_review": got.get("focus_review") or row.focus_review,
            "test_suggestions": got.get("test_suggestions") or row.test_suggestions,
            "risk_description": got.get("risk_description") or row.risk_description,
        }

    tests_by_module: Dict[str, str] = {}
    for idx, row in enumerate(rows):
        got = dict(polished.get(str(idx)) or {})
        item = items[idx] if idx < len(items) else {}
        is_subtag = (row.block_kind or "") == "subtag"
        if is_subtag:
            ref = overall_polished_by_module.get(row.module_name) or {}
            for field in ("modification", "focus_review", "test_suggestions", "risk_description"):
                candidate = str(got.get(field) or "").strip()
                reference = str(ref.get(field) or "").strip()
                draft = str(getattr(row, field, "") or "").strip()
                if candidate and reference and _checklist_text_too_similar(candidate, reference):
                    print(
                        f"[CHECKLIST] 子tag「{row.module_name}」{field} 与总包过相似，回退草稿",
                        flush=True,
                    )
                    got[field] = draft
        # 变更摘要：保留简短范围前缀 + 精炼主题，细节见模块详情链接
        loc_prefix = _normalize_cell_text(str(item.get("file_locations") or ""), max_lines=2)
        if got.get("modification"):
            body = _normalize_cell_text(got["modification"], max_lines=5)
        else:
            body = _normalize_cell_text(row.modification, max_lines=5)
        if loc_prefix and body:
            combined = f"{loc_prefix}\n{body}" if loc_prefix not in body else body
        else:
            combined = loc_prefix or body
        row.modification = _clip(combined, _CLIP_MODIFICATION)
        if got.get("focus_review"):
            row.focus_review = _clip(_normalize_cell_text(got["focus_review"], max_lines=12), _CLIP_FOCUS)
        else:
            row.focus_review = _clip(_normalize_cell_text(row.focus_review, max_lines=12), _CLIP_FOCUS)
        if got.get("test_suggestions"):
            row.test_suggestions = _clip(_normalize_cell_text(got["test_suggestions"], max_lines=4), _CLIP_TEST)
        else:
            row.test_suggestions = _clip(_normalize_cell_text(row.test_suggestions, max_lines=4), _CLIP_TEST)
        if got.get("risk_description"):
            row.risk_description = _clip(_normalize_cell_text(got["risk_description"], max_lines=4), _CLIP_RISK)
        elif row.risk_description:
            row.risk_description = _clip(_normalize_cell_text(row.risk_description, max_lines=4), _CLIP_RISK)
        row.negative_impact = "是" if str(row.negative_impact).strip().startswith("是") else (
            "否" if str(row.negative_impact).strip() else ""
        )
        if row.negative_impact != "是":
            row.risk_description = ""
        row.owner_comment = ""
        # 基础测试页只缓存总包结论，避免被子 tag 行覆盖
        if (row.block_kind or "overall") != "subtag" and row.module_name and not tests_by_module.get(row.module_name):
            tests_by_module[row.module_name] = row.test_suggestions
    task["_checklist_polished_tests"] = tests_by_module


def _checklist_text_too_similar(left: str, right: str) -> bool:
    """Detect when sub-tag polished text is essentially a copy of overall text."""
    a = _normalize_cell_text(left or "")
    b = _normalize_cell_text(right or "")
    if not a or not b:
        return False
    compact_a = re.sub(r"\s+", "", a)
    compact_b = re.sub(r"\s+", "", b)
    if compact_a == compact_b:
        return True
    shorter, longer = (compact_a, compact_b) if len(compact_a) <= len(compact_b) else (compact_b, compact_a)
    if len(shorter) >= 36 and shorter in longer:
        return True
    lines_a = {ln.strip("• ").strip() for ln in a.splitlines() if ln.strip()}
    lines_b = {ln.strip("• ").strip() for ln in b.splitlines() if ln.strip()}
    if not lines_a or not lines_b:
        return False
    inter = len(lines_a & lines_b)
    return inter / max(len(lines_a), len(lines_b)) >= 0.6


def _patch_digest_for_checklist(patches: Sequence[Dict[str, Any]], *, limit: int = 16) -> str:
    """Compact per-patch digest for checklist polish (especially sub-tag rows)."""
    lines: List[str] = []
    for p in patches[:limit]:
        sha = str(p.get("sha") or p.get("commit") or "")[:12]
        title = str(p.get("title") or "").strip()
        brief = str(p.get("problem_solved") or p.get("modification_interpretation") or "").strip()
        if not brief:
            rich = p.get("problem_closure_rich") or {}
            if isinstance(rich, dict):
                brief = str(rich.get("conclusion") or rich.get("why_fixed") or "").strip()
        if not brief:
            rich = p.get("commit_interpretation_rich") or {}
            if isinstance(rich, dict):
                brief = str(rich.get("summary") or "").strip()
        risk = str(p.get("key_risk") or "").strip()
        bits = [x for x in (sha, title, brief[:160] if brief else "", risk[:80] if risk and not is_key_risk_none(risk) else "") if x]
        if bits:
            lines.append(" | ".join(bits))
    return "\n".join(lines)


def _export_test_suggestions(
    task: Dict[str, Any],
    module_name: str,
    analysis: Dict[str, Any],
    patches: Sequence[Dict[str, Any]],
) -> str:
    cached = (task.get("_checklist_polished_tests") or {}).get(module_name)
    if cached:
        return cached
    return _clip(_normalize_cell_text(_module_test_suggestions(analysis, patches)), 360)


def _normalize_cell_text(text: str, *, max_lines: int = 12) -> str:
    text = _plain_text(text)
    lines: List[str] = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        line = re.sub(r"^(?:[-*•]|\d+[.)])\s*", "", line)
        if line:
            lines.append(f"• {line}")
    return "\n".join(lines[:max_lines])


def _module_stats_line(patches: Sequence[Dict[str, Any]]) -> str:
    if not patches:
        return ""
    try:
        from web_app.module_analysis_service import build_module_stats
        stats = build_module_stats(list(patches))
    except Exception:
        return f"共 {len(patches)} 笔 patch"
    return (
        f"共 {stats.get('patch_count', len(patches))} 笔 patch，"
        f"{stats.get('files_changed', 0)} 个文件，"
        f"+{stats.get('insertions', 0)}/-{stats.get('deletions', 0)} 行；"
        f"均分 {stats.get('avg_score', 0)}；"
        f"阻断 {stats.get('blocked_count', 0)} 笔，"
        f"中/高死机风险 {stats.get('elevated_crash_count', 0)} 笔"
    )


def _module_file_locations_summary(
    patches: Sequence[Dict[str, Any]],
    *,
    code_tree: Optional[Dict[str, Any]] = None,
    limit_dirs: int = 12,
    limit_files: int = 8,
) -> str:
    """汇总模块 patch 数量与主要变更目录/文件。"""
    if not patches:
        return ""
    file_hits: Counter[str] = Counter()
    dir_hits: Counter[str] = Counter()
    for p in patches:
        seen: set = set()
        for raw in p.get("files") or []:
            fp = str(raw).strip().lstrip("./")
            if not fp or fp in seen:
                continue
            seen.add(fp)
            file_hits[fp] += 1
            parts = [x for x in fp.split("/") if x]
            if not parts:
                continue
            if len(parts) >= 3:
                dir_hits["/".join(parts[:3]) + "/"] += 1
            if len(parts) >= 2:
                dir_hits["/".join(parts[:2]) + "/"] += 1
            dir_hits[parts[0] + "/"] += 1

    patch_count = len(patches)
    file_count = len(file_hits)
    lines = [f"本模块共 {patch_count} 笔 patch，涉及 {file_count} 个变更文件。"]

    tree_dirs = _code_tree_top_dirs(code_tree, limit=limit_dirs)
    if tree_dirs:
        lines.append(
            "主要修改目录："
            + "；".join(f"{path}({cnt}处变更)" for path, cnt in tree_dirs)
        )
    else:
        candidates = [
            (d, c) for d, c in dir_hits.most_common(limit_dirs * 4)
            if c >= 1 and d.count("/") >= 2
        ]
        top_dirs: List[Tuple[str, int]] = []
        for d, c in candidates:
            if any(d != kept and d.startswith(kept) for kept, _ in top_dirs):
                continue
            top_dirs.append((d, c))
            if len(top_dirs) >= limit_dirs:
                break
        if top_dirs:
            lines.append(
                "主要修改目录："
                + "；".join(f"{d}({c}处变更)" for d, c in top_dirs)
            )

    top_files = file_hits.most_common(limit_files)
    if top_files:
        lines.append(
            "高频变更文件："
            + "；".join(f"{fp}({cnt})" for fp, cnt in top_files)
        )
    return "\n".join(lines)


def _code_tree_top_dirs(
    code_tree: Optional[Dict[str, Any]],
    *,
    limit: int = 12,
    min_depth: int = 2,
    max_depth: int = 4,
) -> List[Tuple[str, int]]:
    """从 module_summaries 的 code_tree 提取 patch 集中的目录路径。"""
    if not code_tree:
        return []
    rows: List[Tuple[str, int]] = []

    def _walk(node: Dict[str, Any], prefix: str, depth: int) -> None:
        name = str(node.get("name") or "")
        if name in ("", "/"):
            path = prefix
        else:
            path = f"{prefix}/{name}".lstrip("/")
        patch_count = int(node.get("patch_count") or 0)
        node_type = str(node.get("type") or "dir")
        if (
            node_type == "dir"
            and path
            and min_depth <= depth <= max_depth
            and patch_count > 0
        ):
            rows.append((path + "/", patch_count))
        for child in node.get("children") or []:
            if isinstance(child, dict):
                _walk(child, path, depth + 1)

    _walk(code_tree, "", 0)
    rows.sort(key=lambda x: (-x[1], x[0]))
    deduped: List[Tuple[str, int]] = []
    seen_prefix: set = set()
    for path, cnt in rows:
        if any(path.startswith(p) and path != p for p, _ in deduped):
            continue
        if path in seen_prefix:
            continue
        seen_prefix.add(path)
        deduped.append((path, cnt))
        if len(deduped) >= limit:
            break
    return deduped


def _analysis_context_for_polish(analysis: Dict[str, Any]) -> str:
    if not analysis:
        return ""
    parts: List[str] = []
    for key, label in (
        ("modification_interpretation", "模块变更解读"),
        ("impact_scope", "影响范围"),
        ("problem_closure_summary", "问题闭环"),
        ("cross_module_notes", "跨模块说明"),
        ("crash_risk_reason", "死机风险依据"),
    ):
        text = str(analysis.get(key) or "").strip()
        if text:
            parts.append(f"{label}：\n{text}")
    changes = analysis.get("code_changes") or []
    if isinstance(changes, list) and changes:
        lines = ["关键代码变化："]
        for item in changes[:12]:
            if not isinstance(item, dict):
                continue
            loc = str(item.get("location") or "").strip()
            exp = str(item.get("explanation") or "").strip()
            risk = str(item.get("risk_level") or "").strip()
            bit = " — ".join(x for x in (loc, exp, f"风险{risk}" if risk else "") if x)
            if bit:
                lines.append(f"- {bit}")
        if len(lines) > 1:
            parts.append("\n".join(lines))
    high_risk = analysis.get("high_risk_patches") or []
    if isinstance(high_risk, list) and high_risk:
        lines = ["高风险 patch："]
        for item in high_risk[:8]:
            if not isinstance(item, dict):
                continue
            sha = str(item.get("sha") or "")[:12]
            title = str(item.get("title") or "")[:80]
            reason = str(item.get("reason") or "")[:120]
            lines.append(f"- {sha} {title} | {reason}".strip(" |"))
        if len(lines) > 1:
            parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _patch_title_themes(patches: Sequence[Dict[str, Any]], limit: int = 10) -> List[Tuple[str, int]]:
    themes: Counter[str] = Counter()
    for p in patches:
        title = str(p.get("title") or "").strip()
        if not title:
            continue
        if ":" in title:
            prefix = title.split(":", 1)[0].strip().lower()
        else:
            prefix = " ".join(title.split()[:3]).lower()[:36]
        if prefix:
            themes[prefix] += 1
    return themes.most_common(limit)


def _overall_module_rows(
    from_tag: str,
    to_tag: str,
    repo: str,
    task: Dict[str, Any],
    grouped: Dict[str, List[Dict[str, Any]]],
    modules_map: Dict[str, Any],
    owners: Dict[str, ModuleOwner],
) -> List[ChecklistRow]:
    tag_text = _overall_tag_text(from_tag, to_tag, repo)
    submit_info = _overall_submit_info(from_tag, to_tag, task)
    names = list(get_dev_module_names())
    extra = [name for name in grouped if name not in names]
    rows: List[ChecklistRow] = []
    for name in names + extra:
        mod_patches = grouped.get(name) or []
        analysis = {}
        entry = modules_map.get(name)
        if isinstance(entry, dict):
            analysis = entry.get("analysis") or {}
        rows.append(
            _build_module_row(
                name,
                mod_patches,
                analysis,
                owners.get(name, ModuleOwner()),
                tag_text=tag_text,
                submit_info=submit_info,
                include_empty=True,
                block_kind="overall",
            )
        )
    return rows


def _subtag_module_rows(
    sub: Dict[str, Any],
    owners: Dict[str, ModuleOwner],
) -> List[ChecklistRow]:
    """子 tag 段只写本段相邻增量代码实际涉及的模块，不多列未改动的模块。

    ``patches`` = rN-1→rN 增量；勿用累计 conclusion_patches 选模块，否则会把
    更早子 tag 已出现过的模块重复铺满后面每一段。
    """
    incremental = list(sub.get("patches") or [])
    if not incremental:
        # 兼容极旧数据：没有增量列表时再退回累计
        incremental = list(sub.get("conclusion_patches") or [])
    grouped = _group_patches(incremental)
    ordered = sort_modules_by_definition(list(grouped.items()))
    rows: List[ChecklistRow] = []
    for name, mod_patches in ordered:
        if not mod_patches:
            continue
        rows.append(
            _build_module_row(
                name,
                mod_patches,
                {},
                owners.get(name, ModuleOwner()),
                tag_text=sub.get("tag_text") or "",
                submit_info="",
                include_empty=False,
                block_kind="subtag",
            )
        )
    return rows


def _build_module_row(
    module_name: str,
    patches: Sequence[Dict[str, Any]],
    analysis: Dict[str, Any],
    owner: ModuleOwner,
    *,
    tag_text: str,
    submit_info: str,
    include_empty: bool,
    block_kind: str = "overall",
) -> ChecklistRow:
    skipped = len(patches) == 0
    if skipped and not include_empty:
        return ChecklistRow(module_name=module_name, skipped=True, block_kind=block_kind)
    domain = f"{module_name}-{owner.owner}" if owner.owner else module_name
    patch_list = list(patches)
    if skipped:
        return ChecklistRow(
            module_name=module_name,
            tag_text=tag_text,
            submit_info=submit_info,
            department=_department_label(owner),
            domain_owner=domain,
            reviewer=owner.reviewer,
            review_plus="不涉及",
            skipped=True,
            block_kind=block_kind,
            source_patches=[],
        )
    return ChecklistRow(
        module_name=module_name,
        tag_text=tag_text,
        submit_info=submit_info,
        department=_department_label(owner),
        domain_owner=domain,
        reviewer=owner.reviewer,
        review_plus="",
        modification=_clip(_module_modification(analysis, patch_list), _DRAFT_MODIFICATION),
        focus_review=_clip(_module_focus_review(analysis, patch_list), _DRAFT_FOCUS),
        negative_impact=_negative_impact_flag(analysis, patch_list),
        risk_description=_clip(_risk_description(analysis, patch_list), 500),
        must_merge=_must_merge(patch_list),
        test_suggestions=_clip(_module_test_suggestions(analysis, patch_list), 500),
        owner_comment="",
        skipped=False,
        block_kind=block_kind,
        source_patches=patch_list,
    )


def _iter_incremental_subtags(
    task: Dict[str, Any],
    patches: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    doc = _load_json_field(task.get("monthly_subtag_results"))
    intervals = list((doc or {}).get("intervals") or [])
    if not intervals:
        return []
    intervals.sort(key=_interval_sort_key)
    task_to = _parse_month_subtag(str(task.get("to_tag") or ""))
    if task_to:
        intervals = [
            item for item in intervals
            if _subtag_not_after_task(item, task_to)
        ]

    all_shas: List[str] = []
    seen_sha: set = set()
    for item in intervals:
        for sha in item.get("patch_shas") or []:
            text = str(sha).strip()
            if text and text not in seen_sha:
                seen_sha.add(text)
                all_shas.append(text)
    extra_rows = get_patches_by_shas(all_shas) if all_shas else []
    pool = list(patches) + [_hydrate_patch(row) for row in extra_rows]
    by_full, by_prefix = _index_patches(pool)

    mode = str((doc or {}).get("analysis_mode") or "")
    adjacent_mode = mode.startswith("adjacent") or _intervals_look_adjacent(intervals)

    prev_shas: set = set()
    prev_r = 0
    blocks: List[Dict[str, Any]] = []
    baseline = str((doc or {}).get("baseline_tag") or "")
    parsed_base = _parse_month_subtag(baseline)
    if parsed_base:
        prev_r = parsed_base[1]
    base_r = parsed_base[1] if parsed_base else 0
    if not base_r and intervals:
        first_from = _parse_month_subtag(str(intervals[0].get("from_tag") or ""))
        if first_from:
            base_r = first_from[1]
    for item in intervals:
        to_tag = str(item.get("to_tag") or "")
        from_tag = str(item.get("from_tag") or "")
        parsed = _parse_month_subtag(to_tag)
        cur_r = parsed[1] if parsed else 0
        parsed_from = _parse_month_subtag(from_tag)
        from_r = parsed_from[1] if parsed_from else 0
        shas = {str(s).strip() for s in (item.get("patch_shas") or []) if str(s).strip()}
        if adjacent_mode:
            incremental = shas
            cumulative_shas = prev_shas | shas
            prev_shas = cumulative_shas
            diff_from = from_r or prev_r or (cur_r - 1 if cur_r else 0)
        else:
            # 兼容旧任务：区间结果已是累计 r1→rN，差分得到相邻增量展示
            incremental = shas - prev_shas
            cumulative_shas = shas
            prev_shas = shas
            diff_from = prev_r or (cur_r - 1 if cur_r else 0)
        if not incremental:
            if cur_r:
                prev_r = cur_r
            continue
        inc_patches = _patches_for_shas(incremental, by_full, by_prefix)
        if not inc_patches:
            if cur_r:
                prev_r = cur_r
            continue
        conclusion_patches = _patches_for_shas(cumulative_shas, by_full, by_prefix) or inc_patches
        sha = _tag_sha(to_tag)
        conclusion_from = base_r or diff_from
        blocks.append({
            "to_tag": to_tag,
            "tag_text": _subtag_cell_text(
                to_tag, sha, diff_from, cur_r, conclusion_from=conclusion_from,
            ),
            "patches": inc_patches,
            "conclusion_patches": conclusion_patches,
        })
        if cur_r:
            prev_r = cur_r
    return blocks


def _intervals_look_adjacent(intervals: Sequence[Dict[str, Any]]) -> bool:
    """Heuristic: consecutive intervals chain as rN-1→rN, rN→rN+1."""
    if len(intervals) < 2:
        # 单区间也可能是 r1→r2；仅当 from 不是统一 baseline 时视为相邻
        if len(intervals) == 1:
            item = intervals[0]
            fr = _parse_month_subtag(str(item.get("from_tag") or ""))
            tr = _parse_month_subtag(str(item.get("to_tag") or ""))
            return bool(fr and tr and tr[1] == fr[1] + 1)
        return False
    ordered = sorted(intervals, key=_interval_sort_key)
    for prev, cur in zip(ordered, ordered[1:]):
        if str(prev.get("to_tag") or "") != str(cur.get("from_tag") or ""):
            return False
    return True


def _overall_tag_text(from_tag: str, to_tag: str, repo: str) -> str:
    lines = [from_tag or to_tag, "", ""]
    if from_tag and to_tag and repo:
        lines.append("diff：")
        lines.append(f"{repo}/+log/{from_tag}..{to_tag}")
    elif from_tag and to_tag:
        lines.append(f"diff：{from_tag}..{to_tag}")
    return "\n".join(lines).rstrip()


def _overall_submit_info(from_tag: str, to_tag: str, task: Dict[str, Any]) -> str:
    lines = [
        f"本次GKI版本升级目标版本：",
        to_tag or "",
        f"升级前GKI版本：{from_tag}" if from_tag else "",
    ]
    return "\n".join(line for line in lines if line).strip()


def _subtag_cell_text(
    to_tag: str,
    sha: str,
    diff_from: int,
    diff_to: int,
    *,
    conclusion_from: int = 0,
) -> str:
    lines = [to_tag]
    if sha:
        lines.append(f"SHA-1: {sha[:20]}")
    if diff_from and diff_to and diff_to != diff_from:
        lines.append(f"Diff: r{diff_from}..r{diff_to}")
    if conclusion_from and diff_to and (conclusion_from != diff_from or conclusion_from != diff_to):
        lines.append(f"结论: r{conclusion_from}..r{diff_to} 累计")
    lines.append("LICENSES")
    return "\n".join(lines)


def _tag_sha(tag: str) -> str:
    try:
        from web_app.analyzer_service import _peel_tag_commit
        return (_peel_tag_commit(tag) or "").strip()
    except Exception:
        return ""


def _module_modification(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    """表格「变更摘要」草稿：只写关键信息，细节留给模块详情页。"""
    parts: List[str] = []
    stats = _module_stats_line(patches)
    if stats:
        parts.append(stats)

    for key in ("problem_closure_summary", "modification_interpretation"):
        text = str(analysis.get(key) or "").strip()
        if text:
            # 取前两句/前两段，避免整段分析塞进表格
            snippet = _shorten_paragraph(text, max_chars=180)
            if snippet and snippet not in parts:
                parts.append(snippet)
            break

    themes = _patch_title_themes(patches)
    if themes:
        theme_text = "；".join(f"{name}({cnt})" for name, cnt in themes[:5])
        parts.append(f"主题：{theme_text}")

    if len(parts) <= 1:
        titles = [str(p.get("title") or "").strip() for p in patches[:4] if p.get("title")]
        if titles:
            parts.append("代表提交：" + "；".join(_unique(titles)[:4]))
    return "\n".join(parts)


def _shorten_paragraph(text: str, *, max_chars: int = 180) -> str:
    cleaned = re.sub(r"\s+", " ", (text or "").strip())
    if not cleaned:
        return ""
    if len(cleaned) <= max_chars:
        return cleaned
    cut = cleaned[:max_chars]
    for sep in ("。", "；", ";", "，", ","):
        idx = cut.rfind(sep)
        if idx >= max_chars // 2:
            return cut[: idx + 1]
    return cut.rstrip() + "…"


def _module_focus_review(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    """按 AI 信号分级：P0 必看 / P1 重点 / 其余可抽样。"""
    total = len(patches)
    buckets = {"P0": [], "P1": [], "P2": []}
    seen: set = set()

    # 模块分析里已标高风险的优先抬到 P0/P1
    high_risk = analysis.get("high_risk_patches") or []
    hr_reason: Dict[str, str] = {}
    if isinstance(high_risk, list):
        for item in high_risk:
            if not isinstance(item, dict):
                continue
            sha = str(item.get("sha") or "").strip().lower()
            if not sha:
                continue
            hr_reason[sha[:12]] = str(item.get("reason") or "").strip()[:80]

    for p in patches:
        level, reason = _patch_review_level(p, hr_reason)
        if level not in buckets:
            continue
        key = str(p.get("sha") or p.get("commit") or id(p))
        if key in seen:
            continue
        seen.add(key)
        line = _format_focus_patch_line(p, reason)
        if line:
            buckets[level].append(line)

    lines: List[str] = []
    p0, p1, p2 = buckets["P0"], buckets["P1"], buckets["P2"]

    if p0:
        lines.append(f"P0 必须逐笔（{len(p0)}）：")
        lines.extend(f"• {x}" for x in p0[:8])
        if len(p0) > 8:
            lines.append(f"• …另有 {len(p0) - 8} 笔同级，见模块详情")
    if p1:
        lines.append(f"P1 建议重点看（{len(p1)}）：")
        lines.extend(f"• {x}" for x in p1[:8])
        if len(p1) > 8:
            lines.append(f"• …另有 {len(p1) - 8} 笔同级，见模块详情")
    if not p0 and not p1:
        if p2:
            lines.append(f"P2 可抽样（{len(p2)}）：")
            lines.extend(f"• {x}" for x in p2[:5])
        else:
            lines.append("未见 P0/P1，可按模块详情抽样核对")
    elif p2 and len(p0) + len(p1) < 6:
        lines.append(f"P2 可抽样（{len(p2)}）：")
        lines.extend(f"• {x}" for x in p2[:3])

    if total > 0 and (p0 or p1):
        lines.append(f"其余约 {max(0, total - len(p0) - len(p1))} 笔可抽样/对照模块详情，不必逐笔")
    elif total > 0 and not (p0 or p1):
        lines.append(f"共 {total} 笔，优先看上方抽样项")

    crash = str(analysis.get("crash_risk") or "").strip()
    crash_reason = str(analysis.get("crash_risk_reason") or "").strip()
    if crash in ("中", "高"):
        lines.insert(0, f"模块死机风险{crash}" + (f"：{crash_reason[:80]}" if crash_reason else ""))
    return "\n".join(lines)


def _patch_review_level(
    patch: Dict[str, Any],
    hr_reason: Optional[Dict[str, str]] = None,
) -> Tuple[str, str]:
    """Return (P0|P1|P2|'', reason). Empty means not listed in focus column."""
    reasons: List[str] = []
    score = float(patch.get("score") or 0)
    rec = str(patch.get("recommendation") or "")
    crash = str(patch.get("crash_risk") or "").strip()
    risk = str(patch.get("key_risk") or "").strip()
    intro = str(patch.get("introduced_risk") or "").strip()
    blocked = bool(patch.get("blocked"))
    sha = str(patch.get("sha") or patch.get("commit") or "").strip().lower()
    sha12 = sha[:12]

    p0 = False
    p1 = False
    if blocked or "不建议" in rec or "阻断" in rec:
        p0 = True
        reasons.append("阻断/不建议合入")
    if crash == "高":
        p0 = True
        reasons.append("死机风险高")
    if score > 0 and score < 4.5:
        p0 = True
        reasons.append(f"低分 {score:.1f}")

    if crash == "中":
        p1 = True
        reasons.append("死机风险中")
    if "有条件合入" in rec:
        p1 = True
        reasons.append("有条件合入")
    if score > 0 and score < 6.5:
        p1 = True
        reasons.append(f"评分 {score:.1f}")
    if risk and not is_key_risk_none(risk):
        p1 = True
        reasons.append("关键风险")
    if intro and intro.lower() not in _NONE_RISK:
        p1 = True
        reasons.append("引入风险")
    if hr_reason and sha12 and sha12 in hr_reason:
        p1 = True
        if hr_reason[sha12]:
            reasons.append(hr_reason[sha12])
        else:
            reasons.append("模块高风险清单")

    # 轻微信号：分数略低或 title 含 fix/CVE 等 → P2
    title = str(patch.get("title") or "").lower()
    p2 = False
    if score > 0 and score < 7.5:
        p2 = True
        reasons.append(f"评分 {score:.1f}")
    if any(k in title for k in ("fix", "cve", "revert", "regression", "null", "uaf", "overflow")):
        p2 = True
        if "关键词" not in "".join(reasons):
            reasons.append("标题敏感词")

    if p0:
        return "P0", "；".join(_unique(reasons)[:3])
    if p1:
        return "P1", "；".join(_unique(reasons)[:3])
    if p2:
        return "P2", "；".join(_unique(reasons)[:2])
    return "", ""


def _format_focus_patch_line(patch: Dict[str, Any], reason: str) -> str:
    sha = str(patch.get("sha") or patch.get("commit") or "")[:12]
    title = str(patch.get("title") or "").strip()[:56]
    bit = " ".join(x for x in (sha, title) if x).strip()
    if not bit:
        return ""
    if reason:
        return f"{bit} — {reason}"
    return bit


def _module_review_comment(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    """兼容旧调用：转调重点 Review 分级文案。"""
    return _module_focus_review(analysis, patches)


def _module_test_suggestions(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    extra: List[str] = []
    checklist = analysis.get("verification_checklist") or []
    if isinstance(checklist, list):
        for item in checklist[:3]:
            if isinstance(item, dict):
                item_text = str(item.get("item") or item.get("text") or "").strip()
            else:
                item_text = str(item).strip()
            if item_text:
                extra.append(item_text)
    if extra:
        return "\n".join(extra)
    text = str(analysis.get("test_suggestions") or "").strip()
    if text:
        return _shorten_paragraph(text, max_chars=240)
    collected: List[str] = []
    for p in patches:
        raw = p.get("llm_test_suggestions") or p.get("test_suggestions")
        rendered = _render_test_suggestions(raw)
        if rendered:
            collected.append(rendered)
    return "\n".join(_unique(collected)[:3])


def _render_test_suggestions(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, str):
        return raw.strip()
    if isinstance(raw, dict):
        lines = []
        for key, value in raw.items():
            if value in (None, "", False):
                continue
            lines.append(f"{key}: {value}")
        return "\n".join(lines)
    return str(raw).strip()


def _negative_impact_flag(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    return "是" if _risk_description(analysis, patches) else "否"


def _risk_description(analysis: Dict[str, Any], patches: Sequence[Dict[str, Any]]) -> str:
    lines: List[str] = []
    crash = str(analysis.get("crash_risk") or "").strip()
    crash_reason = str(analysis.get("crash_risk_reason") or "").strip()
    if crash in ("中", "高"):
        lines.append(f"死机风险{crash}" + (f"：{crash_reason}" if crash_reason else ""))
    elif crash_reason and crash == "低":
        pass
    high_risk = analysis.get("high_risk_patches") or []
    if isinstance(high_risk, list):
        for item in high_risk[:4]:
            if not isinstance(item, dict):
                continue
            title = str(item.get("title") or "").strip()[:50]
            reason = str(item.get("reason") or "").strip()[:80]
            bit = "：".join(x for x in (title, reason) if x)
            if bit:
                lines.append(bit)
    for p in patches:
        if str(p.get("crash_risk") or "").strip() in ("中", "高"):
            reason = str(p.get("crash_risk_reason") or p.get("key_risk") or "").strip()
            title = str(p.get("title") or "")[:40]
            bit = "：".join(x for x in (title, reason) if x)
            if bit:
                lines.append(bit)
        intro = str(p.get("introduced_risk") or "").strip()
        if intro and intro.lower() not in _NONE_RISK:
            lines.append(intro)
        risk = str(p.get("key_risk") or "").strip()
        if risk and not is_key_risk_none(risk):
            lines.append(risk)
    return "\n".join(_unique(lines)[:5])


def _must_merge(patches: Sequence[Dict[str, Any]]) -> str:
    if not patches:
        return "NA"
    recs = [str(p.get("recommendation") or "") for p in patches]
    if any(p.get("blocked") for p in patches):
        return "否"
    if any("不建议合入" in r for r in recs):
        return "否"
    if any("有条件合入" in r or "需大幅修改" in r for r in recs):
        return "NA"
    if any("建议合入" in r for r in recs):
        return "是"
    return "是"


def _group_patches(patches: Sequence[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for p in patches:
        name = str(p.get("dev_module") or "其他").strip() or "其他"
        grouped.setdefault(name, []).append(p)
    return grouped


def _index_patches(patches: Sequence[Dict[str, Any]]) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Dict[str, Any]]]:
    by_full: Dict[str, Dict[str, Any]] = {}
    by_prefix: Dict[str, Dict[str, Any]] = {}
    for p in patches:
        sha = str(p.get("sha") or p.get("commit") or "").strip()
        if not sha:
            continue
        by_full[sha] = p
        by_prefix[sha[:12]] = p
    return by_full, by_prefix


def _patches_for_shas(
    shas: Iterable[str],
    by_full: Dict[str, Dict[str, Any]],
    by_prefix: Dict[str, Dict[str, Any]],
) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    seen = set()
    for sha in shas:
        patch = by_full.get(sha) or by_prefix.get(sha[:12])
        if not patch:
            continue
        ident = str(patch.get("sha") or patch.get("commit") or sha)
        if ident in seen:
            continue
        seen.add(ident)
        out.append(patch)
    return out


def _interval_sort_key(item: Dict[str, Any]) -> Tuple[int, str]:
    to_tag = str(item.get("to_tag") or "")
    parsed = _parse_month_subtag(to_tag)
    return (parsed[1] if parsed else 0, to_tag)


def _subtag_not_after_task(item: Dict[str, Any], task_to: Tuple[str, int]) -> bool:
    parsed = _parse_month_subtag(str(item.get("to_tag") or ""))
    if not parsed:
        return True
    month, r_num = parsed
    if month != task_to[0]:
        return True
    return r_num <= task_to[1]


def _department_label(owner: ModuleOwner) -> str:
    if owner.department and owner.lead:
        return f"{owner.department}@{owner.lead}"
    return owner.department


def _compose_test_strategy(_task: Dict[str, Any]) -> str:
    return _HEADER_PRINCIPLES["测试与验证策略"]


def _load_json_field(raw: Any) -> Dict[str, Any]:
    if not raw:
        return {}
    if isinstance(raw, dict):
        return raw
    try:
        data = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}
    return data if isinstance(data, dict) else {}


def _merge_column(ws: Worksheet, col: int, start: int, end: int) -> None:
    if end <= start:
        return
    ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
    ws.cell(start, col).alignment = _ALIGN_WRAP


def _merge_same_department(
    ws: Worksheet,
    rows: List[ChecklistRow],
    start: int,
    *,
    col: int = _COL_DEPT,
) -> None:
    if not rows:
        return
    run_start = start
    prev = rows[0].department
    for offset, row in enumerate(rows[1:], start=1):
        if row.department != prev:
            _merge_column(ws, col, run_start, start + offset - 1)
            run_start = start + offset
            prev = row.department
    _merge_column(ws, col, run_start, start + len(rows) - 1)


def _row_height(values: Sequence[str], min_h: int = 28, max_h: int = 120) -> int:
    lines = 1
    for value in values:
        text = str(value or "")
        lines = max(lines, text.count("\n") + 1, min(8, (len(text) // 28) + 1))
    return max(min_h, min(max_h, 16 * lines + 8))


def _plain_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"^\s*[-*]\s+", "• ", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _clip(text: str, limit: int = _TEXT_LIMIT) -> str:
    text = _plain_text(text)
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _wrap_cell(text: str, width: int = 36) -> str:
    text = _plain_text(text)
    if not text:
        return ""
    lines: List[str] = []
    for para in text.split("\n"):
        if not para:
            lines.append("")
            continue
        if len(para) <= width or para.startswith("http") or "://" in para[:12]:
            lines.append(para)
            continue
        buf = para
        while len(buf) > width:
            cut = width
            window = buf[: width + 8]
            for sep in "，。；、；;,. ":
                pos = window.rfind(sep)
                if pos >= width * 2 // 3:
                    cut = pos + 1
                    break
            lines.append(buf[:cut].rstrip())
            buf = buf[cut:].lstrip()
        if buf:
            lines.append(buf)
    return "\n".join(lines)


def _wrap_row(values: Sequence[Any], wrap_cols: Dict[int, int]) -> List[Any]:
    out: List[Any] = list(values)
    for col, width in wrap_cols.items():
        idx = col - 1
        if 0 <= idx < len(out) and out[idx] and not isinstance(out[idx], dict):
            out[idx] = _wrap_cell(str(out[idx]), width)
    return out


def _height_from_lines(text: str, min_h: int, max_h: int, line_px: int = 18) -> int:
    lines = max(1, str(text or "").count("\n") + 1)
    return max(min_h, min(max_h, line_px * lines + 12))


def _overview_row_heights(values: List[List[Any]], header_row: int, skip_rows: Sequence[int]) -> Dict[int, int]:
    skip = set(skip_rows)
    heights: Dict[int, int] = {}
    for idx, row in enumerate(values, start=1):
        if idx < header_row:
            # 元信息区按内容撑开，贴近模板行高
            text = _cell_as_text(row[1] if len(row) > 1 else "")
            heights[idx] = _height_from_lines(text, 22, 96, line_px=16)
        elif idx == header_row:
            heights[idx] = 32
        elif idx in skip:
            heights[idx] = 32
        else:
            tallest = max((_cell_as_text(cell) for cell in row[2:]), default="")
            heights[idx] = _height_from_lines(tallest, 48, 140)
    return heights


def _base_row_heights(values: List[List[Any]], skip_rows: Sequence[int]) -> Dict[int, int]:
    skip = set(skip_rows)
    heights: Dict[int, int] = {1: 44}
    for idx, row in enumerate(values[1:], start=2):
        if idx in skip:
            heights[idx] = 32
        else:
            heights[idx] = _height_from_lines(_cell_as_text(row[4] if len(row) > 4 else ""), 48, 120)
    return heights


def _unique(items: Iterable[str]) -> List[str]:
    out: List[str] = []
    seen = set()
    for item in items:
        key = item.strip()
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def _checklist_filename(task: Dict[str, Any]) -> str:
    from web_app.feishu_service import task_run_stamp

    from_tag = str(task.get("from_tag") or "").strip()
    to_tag = str(task.get("to_tag") or "").strip()
    stamp = task_run_stamp(task)
    if from_tag and to_tag:
        raw = f"GKI升级专项checklist_{from_tag}_to_{to_tag}_{stamp}.xlsx"
    else:
        raw = f"GKI升级专项checklist_{stamp}.xlsx" if stamp else "GKI升级专项checklist.xlsx"
    return re.sub(r'[\\/:*?"<>|]+', "_", raw)
