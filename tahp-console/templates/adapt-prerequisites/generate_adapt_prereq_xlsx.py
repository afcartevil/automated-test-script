# -*- coding: utf-8 -*-
"""生成 AI适配调试 前置信息收集 Excel 模板（与 Console 域字段 key 对齐）。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "AI适配调试_前置信息收集模板.xlsx"

# workspace_id | domain_id | domain_label | related_skills | fields: (key, label, type, required, example, options)
# options: None or list[str] for select
DOMAINS: list[dict] = [
    {
        "workspace_id": "adapt-debug",
        "domain_id": "tp",
        "domain_label": "TP触摸屏",
        "module": "kernel-driver",
        "related_skills": "tp-adapt-bringup",
        "fields": [
            ("debug_project", "调试项目名", "text", True, "X6876_H786", None),
            ("project_platform", "项目平台", "text", True, "MT6878", None),
            ("ic_vendor", "IC厂商", "text", True, "FocalTech / Goodix", None),
            ("ic_model", "IC型号", "text", True, "FT8006M", None),
            ("spi_gpio", "SPI口及GPIO号", "text", True, "SPI0 CSIB 等", None),
            ("tp_int_rst_gpio", "TP int、rst gpio脚", "text", True, "INT=xx RST=yy", None),
            ("vddi_avdd_power", "vddi、avdd的LDO或GPIO", "text", True, "LDO / GPIO 配置", None),
            ("old_tp_firmware", "升级前TP固件文件", "text", True, "绝对路径", None),
            ("old_dts", "升级前dts文件", "text", True, "绝对路径", None),
            ("ref_project_branch", "同分支参考项目", "text", True, "参考项目名或路径", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "vib",
        "domain_label": "振动马达",
        "module": "kernel-driver",
        "related_skills": "vib-adapt-lra-driver",
        "fields": [
            ("motor_type", "马达类型", "text", True, "LRA / ERM", None),
            ("ic_vendor", "控制IC厂商", "text", True, "Awinic / AAC", None),
            ("ic_model", "控制IC型号", "text", True, "AW8692x", None),
            ("motor_model", "马达型号", "text", True, "马达型号", None),
            ("i2c_bus", "i2c号", "text", True, "7", None),
            ("i2c_addr", "i2c通讯地址", "text", True, "0x5A", None),
            ("irq_gpio", "马达中断gpio", "text", True, "IRQ GPIO", None),
            ("rst_gpio", "复位gpio", "text", True, "RST GPIO", None),
            ("old_motor_firmware", "升级前的马达固件", "text", True, "绝对路径", None),
            ("old_dts_config", "升级前的dts参数配置", "text", True, "绝对路径", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "lcm",
        "domain_label": "LCM",
        "module": "kernel-driver",
        "related_skills": "display_analyser",
        "fields": [
            ("PROJECT_NAME", "PROJECT_NAME（项目名称）", "text", True, "项目名称", None),
            ("PLATFORM_NAME", "PLATFORM_NAME（平台名称）", "text", True, "平台名称", None),
            ("BOARD_NAME", "BOARD_NAME（主板名称）", "text", True, "主板名称", None),
            ("PANEL_TYPE", "PANEL_TYPE（屏幕类型）", "text", True, "屏幕类型", None),
            ("KERNEL_VERSION_OLD", "KERNEL_VERSION_OLD（旧项目内核版本）", "text", True, "kernel-5.15", None),
            ("KERNEL_VERSION_NEW", "KERNEL_VERSION_NEW（新项目内核版本）", "text", True, "kernel-6.1", None),
            ("PROJ_FILE_OLD", "PROJ_FILE_OLD（旧项目代码目录）", "text", True, "绝对路径", None),
            ("PROJ_FILE_NEW", "PROJ_FILE_NEW（新项目代码目录）", "text", True, "绝对路径", None),
            ("AUTHOR_NAME", "AUTHOR_NAME（作者姓名）", "text", True, "姓名", None),
            ("DATE", "DATE（日期）", "text", True, "2026-07-15", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "fp",
        "domain_label": "指纹",
        "module": "kernel-driver",
        "related_skills": "drv-fp-adapt",
        "fields": [
            ("debug_project", "调试项目名", "text", True, "LJ8k", None),
            ("project_platform", "项目平台", "text", True, "MT6878", None),
            ("kernel_version", "kernel版本", "text", True, "kernel-6.1", None),
            ("fp_vendor", "指纹厂商", "text", True, "goodix", None),
            ("fp_ic_model", "指纹IC型号", "text", True, "GW9598", None),
            ("code_path_old", "升级前代码仓路径", "text", True, "绝对路径", None),
            ("code_path_new", "升级后代码仓路径", "text", True, "绝对路径", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "audio",
        "domain_label": "音频",
        "module": "kernel-driver",
        "related_skills": "audio-feature-adaptation",
        "fields": [
            ("vnd_path_old", "升级前vnd绝对路径", "text", True, "绝对路径", None),
            ("vnd_path_new", "升级后vnd绝对路径", "text", True, "绝对路径", None),
            ("upgrade_scenario", "升级场景", "text", True, "如 芯片升级/分支升级", None),
            ("old_build_option", "升级前编译选项", "text", True, "编译 option", None),
            ("sw_config_screenshot", "软件配置表截图/说明", "textarea", True, "可贴路径或文字说明", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "charging",
        "domain_label": "充电",
        "module": "kernel-driver",
        "related_skills": "mtk-adapt-upgrade-branch-orchestrator",
        "fields": [
            ("mode", "适配类型", "select", True, "upgrade", ["upgrade", "greenfield", "partial"]),
            ("project_full", "项目全名", "text", True, "完整项目名", None),
            ("aosp_root", "新分支代码根目录", "text", True, "绝对路径", None),
            ("platform", "芯片平台", "text", False, "MTK / QCOM / UNISOC", None),
            ("kernel_version", "新分支内核版本号", "text", False, "6.1", None),
            ("old_tran_root", "老分支tran-projects根目录", "text", True, "绝对路径", None),
            ("pillar_required", "适配哪些功能", "select", True, "1&2", ["1", "2", "1&2"]),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "sensor",
        "domain_label": "Sensor",
        "module": "kernel-driver",
        "related_skills": "sensor-feature-adaption",
        "fields": [
            ("old_project_path", "旧项目完整路径", "text", True, "绝对路径", None),
            ("new_baseline_path", "新基线完整路径", "text", True, "绝对路径", None),
            ("project_name", "项目名", "text", True, "项目名", None),
            ("new_features", "具体新增特性", "textarea", True, "列出新增特性", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "security",
        "domain_label": "安全",
        "module": "kernel-driver",
        "related_skills": "android-platform-security-upgrade-ai-debug",
        "fields": [
            ("branch_path_old", "升级前分支路径", "text", True, "绝对路径", None),
            ("branch_path_new", "升级后分支路径", "text", True, "绝对路径", None),
            ("project_name", "项目名", "text", True, "项目名", None),
            ("platform_name", "平台名", "text", True, "平台名", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "mini-system",
        "domain_label": "小系统",
        "module": "kernel-driver",
        "related_skills": "KRN_LK_CPU_Ddapt",
        "fields": [
            ("lk2_path_old", "升级前lk2路径", "text", True, "绝对路径", None),
            ("lk2_path_new", "升级后lk2路径", "text", True, "绝对路径", None),
            ("project_name", "项目名", "text", True, "项目名", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "modem",
        "domain_label": "通信Modem",
        "module": "telephony",
        "related_skills": "MODEM_P1_UPGRADE_SKILL",
        "fields": [
            ("path_a_pub", "Path_A_Pub（原分支公共仓）", "text", True, "绝对路径", None),
            ("path_a_priv", "Path_A_Priv（原分支私有仓）", "text", True, "绝对路径", None),
            ("path_b_pub", "Path_B_Pub（目标分支公共仓）", "text", True, "绝对路径", None),
            ("path_b_priv", "Path_B_Priv（目标分支私有仓）", "text", True, "绝对路径", None),
            ("bom_name", "BOM_Name（项目 BOM）", "text", True, "BOM 名", None),
            ("bom_name_target", "BOM_Name_Target（变更后 BOM）", "text", True, "目标 BOM", None),
            ("rf_target_location", "RF_TARGET_LOCATION", "select", True, "public", ["public", "private"]),
            ("modem_platform", "Modem 平台（补充）", "select", False, "mtk", ["mtk", "qcom", "unisoc"]),
            ("rf_param_dir", "RF 参数目录（补充）", "text", False, "可选", None),
            ("carrier", "运营商（补充）", "text", False, "可选", None),
            ("user_notes", "补充说明", "textarea", False, "可选", None),
            ("AP_Project_Path", "AP_Project_Path（Skill补充，Console暂无）", "text", False, "可选预填", None),
            ("GPIO_Path_C", "GPIO_Path_C（Skill补充）", "text", False, "可选预填", None),
            ("GPIO_Path_D", "GPIO_Path_D（Skill补充）", "text", False, "可选预填", None),
            ("Is_TOS16_Before", "Is_TOS16_Before（Skill补充）", "text", False, "yes/no", None),
        ],
    },
    {
        "workspace_id": "adapt-debug",
        "domain_id": "comm-framework",
        "domain_label": "通信框架",
        "module": "telephony",
        "related_skills": "telephony-version-adapter",
        "fields": [
            ("prj_path", "PRJ Path（项目路径）", "text", True, "绝对路径", None),
            ("user_notes", "补充说明", "textarea", False, "可选", None),
        ],
    },
    # 扩展工作区：特性适配（摘要域，便于同仓收集）
    {
        "workspace_id": "feature-adapt",
        "domain_id": "tp",
        "domain_label": "TP触摸屏(特性适配)",
        "module": "kernel-driver",
        "related_skills": "tp-adapt-bringup",
        "fields": [
            ("repo_path", "代码仓库路径", "text", True, "绝对路径", None),
            ("feature_checklist", "特性清单(名称\\t编码，多行)", "textarea", False, "特性名\\tSF-xxxx", None),
            ("user_notes", "补充说明", "text", False, "可选", None),
        ],
    },
    {
        "workspace_id": "carrier-config",
        "domain_id": "ecc-config",
        "domain_label": "ECC紧急号码",
        "module": "telephony",
        "related_skills": "ecc-config",
        "fields": [
            ("repo_path_mtk", "MTK 代码仓路径", "text", False, "三平台至少填一", None),
            ("repo_path_unisoc", "展锐代码仓路径", "text", False, "三平台至少填一", None),
            ("repo_path_qualcomm", "高通代码仓路径", "text", False, "三平台至少填一", None),
            ("jira_key", "Jira 单号", "text", False, "如 TECNO-12345", None),
            ("ecc_entries_json", "ECC条目JSON(可选整表)", "textarea", False, '[{"number":"112",...}]', None),
        ],
    },
    {
        "workspace_id": "carrier-config",
        "domain_id": "apn-config",
        "domain_label": "APN配置",
        "module": "telephony",
        "related_skills": "apn-edit",
        "fields": [
            ("repo_path", "代码仓库路径", "text", True, "绝对路径", None),
            ("extra_info", "补充信息", "textarea", False, "可选", None),
        ],
    },
]


def style_header(cell, fill_color="1F4E79"):
    cell.font = Font(bold=True, color="FFFFFF", name="Microsoft YaHei")
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def style_cell(cell, wrap=True):
    cell.font = Font(name="Microsoft YaHei", size=10)
    cell.alignment = Alignment(wrap_text=wrap, vertical="center")


def thin_border(ws, rows, cols):
    border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(r, c).border = border


def autosize(ws, min_w=12, max_w=42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            length = max(length, min(len(v), max_w))
        ws.column_dimensions[letter].width = max(min_w, length + 2)


def add_readme(wb: Workbook):
    ws = wb.create_sheet("00_使用说明", 0)
    lines = [
        ("AI适配调试 · 前置信息收集模板", True),
        ("", False),
        ("一、用途", True),
        ("按业务域收集启动「AI适配调试」前的必填/可选前置信息；后续可据此自动回填 Console 表单 form_data。", False),
        ("字段名（field_key）与 tahp-console 内置域定义一致（adapt-debug-domains 等）。", False),
        ("", False),
        ("二、怎么填", True),
        ("1. 先在「01_项目登记」新建一行：填写 project_code（唯一）、项目名、责任人、需要适配的业务域列表。", False),
        ("2. 再按域打开对应工作表（如 AD_tp、AD_charging），用同一 project_code 填一行 value。", False),
        ("3. 黄色列为必填 value 区；灰色为示例/说明，勿改表头 field_key 行。", False),
        ("4. select 字段请填选项「值」本身（如 upgrade / public / mtk），不要填中文释义。", False),
        ("5. 路径类请填本机或构建机上的绝对路径，便于 Agent 直接使用。", False),
        ("", False),
        ("三、工作表命名规则", True),
        ("AD_* = workspace adapt-debug（AI适配调试主工作区）", False),
        ("FA_* = workspace feature-adapt（特性适配）", False),
        ("CC_* = workspace carrier-config（运营商配置）", False),
        ("99_字段字典 = 全量字段目录，供导入程序解析", False),
        ("02_ECC条目 = ECC 多号码明细（可选，与 CC_ecc-config 配合）", False),
        ("", False),
        ("四、后续自动回填约定（建议）", True),
        ("读取：project_code + workspace_id + domain_id → 组装 form_data[field_key]=value", False),
        ("跳过空值；required=Y 且值为空 → 导入前校验失败", False),
        ("充电 mode/pillar_required、modem rf_target_location 等保持枚举原值", False),
        ("", False),
        ("五、版本", True),
        ("与 Console 内置域对齐日期：2026-07-15；域布局若被用户自定义，以运行时 manifest 为准并可扩展字段字典。", False),
    ]
    ws.column_dimensions["A"].width = 100
    for i, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(i, 1, text)
        cell.font = Font(name="Microsoft YaHei", size=12 if bold else 11, bold=bold, color="1F4E79" if bold else "24292F")
        cell.alignment = Alignment(wrap_text=True)


def add_project_register(wb: Workbook):
    ws = wb.create_sheet("01_项目登记", 1)
    headers = [
        "project_code",
        "project_name",
        "workspace_ids",
        "domain_ids",
        "owner",
        "collect_date",
        "platform",
        "branch_or_baseline",
        "related_skills_override",
        "notes",
    ]
    labels = [
        "项目编码(唯一)",
        "项目名称",
        "工作区(逗号分隔)",
        "业务域id(逗号分隔)",
        "责任人",
        "收集日期",
        "芯片平台",
        "目标分支/基线说明",
        "关联Skill覆盖(可选)",
        "备注",
    ]
    examples = [
        "PRJ-2026-001",
        "X6878 示例机",
        "adapt-debug",
        "tp,charging,modem",
        "张三",
        "2026-07-15",
        "MT6878",
        "androidU → androidV",
        "",
        "先收 TP/充电/Modem",
    ]
    tips = [
        "主键，与各域表关联",
        "可读名称",
        "adapt-debug / feature-adapt / carrier-config",
        "与 Sheet 名中 domain_id 对应，如 tp,charging",
        "",
        "YYYY-MM-DD",
        "MTK/QCOM/UNISOC 型号",
        "文字说明即可",
        "一般留空，用域默认 Skill",
        "",
    ]
    for c, h in enumerate(headers, 1):
        style_header(ws.cell(1, c, h))
        style_header(ws.cell(2, c, labels[c - 1]), "2E75B6")
        cell = ws.cell(3, c, tips[c - 1])
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        style_cell(cell)
        ex = ws.cell(4, c, examples[c - 1])
        ex.fill = PatternFill("solid", fgColor="FFF2CC")
        style_cell(ex)
    # blank editable rows
    yellow = PatternFill("solid", fgColor="FFFDE7")
    for r in range(5, 15):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c, "")
            cell.fill = yellow
            style_cell(cell)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    thin_border(ws, 14, len(headers))
    autosize(ws)
    ws.freeze_panes = "A5"
    # helper note
    ws.cell(16, 1, "说明：第1行=field_key（程序用）；第2行=中文表头；第3行=填写提示；第4行=示例（可删）；第5行起=正式数据（黄底）。")
    ws.merge_cells(start_row=16, start_column=1, end_row=16, end_column=6)


def add_domain_sheet(wb: Workbook, d: dict):
    prefix = {"adapt-debug": "AD", "feature-adapt": "FA", "carrier-config": "CC"}[d["workspace_id"]]
    title = f"{prefix}_{d['domain_id']}"
    # sheet name max 31
    ws = wb.create_sheet(title[:31])

    # meta block
    meta = [
        ("workspace_id", d["workspace_id"]),
        ("domain_id", d["domain_id"]),
        ("domain_label", d["domain_label"]),
        ("module", d["module"]),
        ("related_skills_default", d["related_skills"]),
    ]
    ws.cell(1, 1, "META").font = Font(bold=True, color="FFFFFF", name="Microsoft YaHei")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="833C0C")
    for i, (k, v) in enumerate(meta, 2):
        ws.cell(i, 1, k).fill = PatternFill("solid", fgColor="FCE4D6")
        ws.cell(i, 2, v)
        style_cell(ws.cell(i, 1))
        style_cell(ws.cell(i, 2))

    # data table starts at row 8
    start = 8
    # Row1 keys, Row2 labels, Row3 required, Row4 type, Row5 example, Row6+ values
    common = [
        ("project_code", "项目编码", "Y", "text", "与01_项目登记一致"),
    ]
    fields = common + [
        (f[0], f[1], "Y" if f[3] else "N", f[2], f[4] or "") for f in d["fields"]
    ]
    fields.append(("collector", "填写人", "N", "text", ""))
    fields.append(("updated_at", "更新日期", "N", "text", "YYYY-MM-DD"))

    for c, (key, label, req, typ, example) in enumerate(fields, 1):
        h = ws.cell(start, c, key)
        style_header(h, "1F4E79" if req == "Y" else "5B9BD5")
        lab = ws.cell(start + 1, c, label)
        style_header(lab, "2E75B6" if req == "Y" else "9DC3E6")
        rcell = ws.cell(start + 2, c, req)
        rcell.fill = PatternFill("solid", fgColor="F4B183" if req == "Y" else "E2EFDA")
        style_cell(rcell)
        tcell = ws.cell(start + 3, c, typ)
        tcell.fill = PatternFill("solid", fgColor="F2F2F2")
        style_cell(tcell)
        ecell = ws.cell(start + 4, c, example)
        ecell.fill = PatternFill("solid", fgColor="DEEBF7")
        style_cell(ecell)

    yellow = PatternFill("solid", fgColor="FFFDE7")
    for r in range(start + 5, start + 12):
        for c in range(1, len(fields) + 1):
            cell = ws.cell(r, c, "")
            cell.fill = yellow
            style_cell(cell)

    # sample first data row: project_code only
    ws.cell(start + 5, 1, "PRJ-2026-001")

    # data validations for select fields
    for c, f in enumerate(d["fields"], 2):  # +1 for project_code
        opts = f[5]
        if opts:
            # openpyxl dv needs comma-separated, quote carefully
            formula = '"' + ",".join(opts) + '"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            dv.error = "请选择列表中的值"
            dv.errorTitle = "非法选项"
            col = get_column_letter(c)
            dv.add(f"{col}{start + 5}:{col}{start + 50}")
            ws.add_data_validation(dv)

    ws.cell(start + 13, 1, "表头含义：第8行 field_key｜第9行 中文名｜第10行 必填Y/N｜第11行 类型｜第12行 示例｜第13行起填写（黄底）")
    ws.merge_cells(start_row=start + 13, start_column=1, end_row=start + 13, end_column=min(6, len(fields)))

    thin_border(ws, start + 11, len(fields))
    autosize(ws, min_w=14, max_w=36)
    ws.freeze_panes = f"B{start + 5}"
    ws.row_dimensions[start + 1].height = 30


def add_ecc_entries(wb: Workbook):
    ws = wb.create_sheet("02_ECC条目")
    headers = [
        "project_code",
        "group_index",
        "number",
        "with_sim",
        "without_sim",
        "display_name",
        "block",
        "user_notes",
    ]
    labels = [
        "项目编码",
        "组序号(从1)",
        "号码",
        "有卡情况",
        "无卡情况",
        "显示名称",
        "是否屏蔽",
        "补充",
    ]
    tips = [
        "关联01_项目登记",
        "1,2,3…",
        "如 112",
        "真紧急/假紧急/非紧急",
        "真紧急/非紧急",
        "可选",
        "yes/no",
        "可选",
    ]
    for c, h in enumerate(headers, 1):
        style_header(ws.cell(1, c, h))
        style_header(ws.cell(2, c, labels[c - 1]), "2E75B6")
        tip = ws.cell(3, c, tips[c - 1])
        tip.fill = PatternFill("solid", fgColor="F2F2F2")
        style_cell(tip)
    yellow = PatternFill("solid", fgColor="FFFDE7")
    for r in range(4, 20):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c, "")
            cell.fill = yellow
            style_cell(cell)
    # sample
    samples = ["PRJ-2026-001", 1, "112", "真紧急", "真紧急", "紧急电话", "no", ""]
    for c, v in enumerate(samples, 1):
        ws.cell(4, c, v)
    for col, opts in [(4, "真紧急,假紧急,非紧急"), (5, "真紧急,非紧急"), (7, "yes,no")]:
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        letter = get_column_letter(col)
        dv.add(f"{letter}4:{letter}100")
        ws.add_data_validation(dv)
    thin_border(ws, 19, len(headers))
    autosize(ws)
    ws.freeze_panes = "A4"
    ws.cell(21, 1, "导入时按 project_code 聚合为 ecc_entries_json，并配合 CC_ecc-config 的仓路径字段。")


def add_dictionary(wb: Workbook):
    ws = wb.create_sheet("99_字段字典")
    headers = [
        "workspace_id",
        "domain_id",
        "domain_label",
        "module",
        "related_skills_default",
        "field_key",
        "field_label",
        "field_type",
        "required",
        "example",
        "select_options",
        "sheet_name",
        "notes",
    ]
    for c, h in enumerate(headers, 1):
        style_header(ws.cell(1, c, h), "375623")
    r = 2
    for d in DOMAINS:
        prefix = {"adapt-debug": "AD", "feature-adapt": "FA", "carrier-config": "CC"}[d["workspace_id"]]
        sheet = f"{prefix}_{d['domain_id']}"[:31]
        for f in d["fields"]:
            key, label, typ, req, example, opts = f
            row = [
                d["workspace_id"],
                d["domain_id"],
                d["domain_label"],
                d["module"],
                d["related_skills"],
                key,
                label,
                typ,
                "Y" if req else "N",
                example or "",
                ",".join(opts) if opts else "",
                sheet,
                "ECC三选一仓路径" if d["domain_id"] == "ecc-config" and key.startswith("repo_path_") else "",
            ]
            for c, v in enumerate(row, 1):
                cell = ws.cell(r, c, v)
                style_cell(cell)
                if req:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
            r += 1
    thin_border(ws, r - 1, len(headers))
    autosize(ws, min_w=10, max_w=40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r - 1}"


def main():
    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)
    add_readme(wb)
    add_project_register(wb)
    add_ecc_entries(wb)
    for d in DOMAINS:
        add_domain_sheet(wb, d)
    add_dictionary(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
