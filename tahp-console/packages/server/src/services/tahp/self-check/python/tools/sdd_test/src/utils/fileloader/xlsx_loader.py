import logging
import re
import zipfile
from os import path
from xml.etree import ElementTree as ET

XLSX_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
XLSX_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _get_first_worksheet_path(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    first_sheet = workbook.find(f".//{XLSX_MAIN_NS}sheet")
    if first_sheet is None:
        return "xl/worksheets/sheet1.xml"

    relation_id = first_sheet.attrib.get(f"{XLSX_REL_NS}id")
    if not relation_id:
        return "xl/worksheets/sheet1.xml"

    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for relation in rels:
        if relation.attrib.get("Id") != relation_id:
            continue
        target = relation.attrib.get("Target", "worksheets/sheet1.xml")
        return path.normpath(path.join("xl", target)).replace("\\", "/")

    return "xl/worksheets/sheet1.xml"


def _load_shared_strings(archive):
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []

    shared_strings = []
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    for shared_item in root.findall(f"{XLSX_MAIN_NS}si"):
        texts = [
            text.text or "" for text in shared_item.findall(f".//{XLSX_MAIN_NS}t")
        ]
        shared_strings.append("".join(texts))
    return shared_strings


def _cell_column_index(cell_ref):
    match = re.match(r"([A-Z]+)", cell_ref or "")
    if not match:
        return 0

    index = 0
    for char in match.group(1):
        index = index * 26 + ord(char) - ord("A") + 1
    return index - 1


def _read_cell_value(cell, shared_strings):
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(
            text.text or "" for text in cell.findall(f".//{XLSX_MAIN_NS}t")
        ).strip()

    value_node = cell.find(f"{XLSX_MAIN_NS}v")
    if value_node is None or value_node.text is None:
        return ""

    value = value_node.text
    if cell_type == "s":
        return shared_strings[int(value)].strip()
    return value.strip()


def _rows_to_dict(rows):
    if not rows:
        return {}

    titles = [str(title).strip() if title is not None else "" for title in rows[0]]
    feature_dict = {}
    for row in rows[1:]:
        if not row:
            continue
        feature_code = str(row[0]).strip() if row[0] is not None else ""
        if not feature_code:
            continue
        feature_dict[feature_code] = {
            title: str(row[index]).strip()
            if index < len(row) and row[index] is not None
            else ""
            for index, title in enumerate(titles[1:], start=1)
            if title
        }

    return feature_dict


def _parse_xlsx_with_openpyxl(xlsx_file):
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_file, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        return _rows_to_dict(rows)
    finally:
        workbook.close()


def _parse_xlsx_with_xml(xlsx_file):
    rows = []
    with zipfile.ZipFile(xlsx_file) as archive:
        shared_strings = _load_shared_strings(archive)
        worksheet_path = _get_first_worksheet_path(archive)
        worksheet = ET.fromstring(archive.read(worksheet_path))

        for row in worksheet.findall(f".//{XLSX_MAIN_NS}sheetData/{XLSX_MAIN_NS}row"):
            row_values = []
            for cell in row.findall(f"{XLSX_MAIN_NS}c"):
                column_index = _cell_column_index(cell.attrib.get("r", ""))
                while len(row_values) <= column_index:
                    row_values.append("")
                row_values[column_index] = _read_cell_value(cell, shared_strings)
            rows.append(row_values)

    return _rows_to_dict(rows)


def parse_xlsx_to_dict(xlsx_file):
    """将xlsx第一列作为key，第一行作为title解析为字典。"""
    if not path.exists(xlsx_file):
        logging.warning(f"xlsx file not found: {xlsx_file}")
        return {}

    try:
        return _parse_xlsx_with_openpyxl(xlsx_file)
    except ImportError:
        logging.info("openpyxl is not installed, parse xlsx with xml fallback")
        return _parse_xlsx_with_xml(xlsx_file)
