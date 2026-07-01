from dataclasses import dataclass, field, fields
from typing import Any, List, Union, Optional, Tuple
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


def title_meta(
        *,
        default: Any = None,
        display_name: Optional[str] = None,
        width: Optional[float] = None,
        bg_color: Optional[str] = None,
        font_color: Optional[str] = None,
        header_align: Tuple[str, str] = ('center', 'center'),
        data_align: Tuple[str, str] = ('left', 'center'),
        number_format: Optional[str] = None,
) -> Any:
    """
    Factory for dataclass fields to attach Excel header & data metadata.

    Non-title_meta fields will have empty metadata and be skipped in writing.

    :param default: default cell value
    :param display_name: header text (defaults to field name)
    :param width: column width
    :param bg_color: header background fill color
    :param font_color: header font color
    :param header_align: (horizontal, vertical) for header cell alignment
    :param data_align: (horizontal, vertical) for data cell alignment
    :param number_format: Excel number format string for data cells

    Usage:
        @dataclass
        class Unit:
            col1: str = title_meta(
                default='', display_name='Column 1', width=20,
                bg_color='FFCCFF', font_color='000000',
                header_align=('center','center'), data_align=('left','center'),
                number_format=None
            )
    """
    return field(
        default=default,
        metadata={
            'display_name': display_name,
            'width': width,
            'bg_color': bg_color,
            'font_color': font_color,
            'header_align': header_align,
            'data_align': data_align,
            'number_format': number_format,
        }
    )


@dataclass
class ExcelTableUnit:
    """
    Base class for Excel table units.
    Extend this class with your own fields, using title_meta to specify header & data settings.
    Non-title_meta fields are ignored when generating Excel output.
    Example:
        @dataclass
        class MyUnit(ExcelTableUnit):
            name: str = title_meta(
                default='', display_name='Name', width=25,
                bg_color='FFEB9C', font_color='000000',
                header_align=('center','center'), data_align=('left','center')
            )
            score: float = title_meta(
                default=0.0, display_name='Score', width=15,
                bg_color='C6EFCE', font_color='006100',
                number_format='0.00', data_align=('right','center')
            )
            remarks: str = title_meta(default='OK')
            # intermediate variable, not included in Excel:
            temp_cache: dict = field(default_factory=dict)
    """
    def _preprocess_source(self, source: dict):
        """
        将Source进行预处理，方便在后续parse中使用
        一般来说，你如果处理了source，应返回一个副本，防止直接影响来源
        """
        return source

    @classmethod
    def parse_from(cls, source: dict):
        ins = cls()
        ins.parse(source)
        return ins

    def parse(self, source: dict):
        _final_source = self._preprocess_source(source)
        for k, v in _final_source.items():
            if hasattr(self, k):
                if not self._on_parse_field(k, v):
                    setattr(self, k, v)
        self._on_parse_finished(_final_source)

    def _on_parse_field(self, filed_name: str, source_value: Any) -> bool:
        return False

    def _on_parse_finished(self, final_source: dict):
        pass


class ExcelTableWriter:
    """
    Write lists of dataclass-based units to sheets in an Excel workbook.
    """

    def __init__(self, excel_filepath: str):
        self.excel_filepath = excel_filepath
        self.workbook = self._create_or_load_workbook(excel_filepath)
        # default styles
        self._default_fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
        self._default_font = Font(name='微软雅黑', bold=True)
        self._default_border = Border(*(Side(style='thin'),) * 4)

    @staticmethod
    def _create_or_load_workbook(path: str) -> openpyxl.Workbook:
        try:
            wb = openpyxl.load_workbook(path)
        except FileNotFoundError:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
        return wb

    def save(self) -> None:
        """Save workbook to disk."""
        self.workbook.save(self.excel_filepath)

    def write_sheet(self, sheet_name: str, units: List[ExcelTableUnit], save_workbook=True) -> None:
        # remove existing sheet
        if sheet_name in self.workbook.sheetnames:
            self.workbook.remove(self.workbook[sheet_name])
        if not units:
            if save_workbook:
                self.save()
            return

        ws = self.workbook.create_sheet(title=sheet_name)
        sample = units[0]
        cols = []  # (name, metadata...)
        for f in fields(sample):
            # skip fields without title_meta metadata
            md = f.metadata
            if not md:
                continue
            header = md.get('display_name') or f.name
            width = md.get('width')
            bg = md.get('bg_color')
            fg = md.get('font_color')
            head_align = md.get('header_align', ('center', 'center'))
            data_align = md.get('data_align', ('left', 'center'))
            num_fmt = md.get('number_format')
            # 如果未指定宽度，则根据字符长度推断（中文按2，其他按1），再加2
            if width is None:
                total = 0
                for ch in header:
                    total += 2 if '\u4e00' <= ch <= '\u9fff' else 1
                width = total + 2
            cols.append((f.name, header, width, bg, fg, head_align, data_align, num_fmt))

        # write header
        for idx, (_, header, width, bg, fg, head_align, _, _) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=idx, value=header)
            if width:
                ws.column_dimensions[cell.column_letter].width = width
            # apply header styles
            cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type='solid') if bg else self._default_fill
            cell.font = Font(name='微软雅黑', bold=True, color=fg) if fg else self._default_font
            cell.alignment = Alignment(horizontal=head_align[0], vertical=head_align[1])
            cell.border = self._default_border

        # header row height & freeze
        ws.row_dimensions[1].height = (ws.row_dimensions[1].height or 15) * 1.5
        ws.freeze_panes = 'A2'

        # write data rows
        for r, unit in enumerate(units, start=2):
            for c, (field_name, *_, data_align, num_fmt) in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=c, value=getattr(unit, field_name))
                # apply data styles
                cell.alignment = Alignment(horizontal=data_align[0], vertical=data_align[1])
                if num_fmt:
                    cell.number_format = num_fmt
                cell.border = self._default_border

        if save_workbook:
            self.save()
