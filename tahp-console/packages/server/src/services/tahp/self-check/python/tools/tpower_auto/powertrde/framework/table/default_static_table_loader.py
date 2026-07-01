import json
import os.path
from typing import Optional

import openpyxl

from powertrde import env
from powertrde.framework.decorators import overwrite
from powertrde.framework.table.base_table_loader import BaseTableLoader


class DefaultStaticTableLoader(BaseTableLoader):
    def __init__(self, table_home: Optional[str] = None, secondary_table_home: Optional[str] = None):
        if table_home is None:
            table_home = env.GET_TEAM_TABLES_DIR()
        if secondary_table_home is None:
            secondary_table_home = env.ASSETS_DIR
        self.table_home = table_home
        self.secondary_table_home = secondary_table_home

    @overwrite
    def load_table(self, table_name: str, primary_key_to_lower=False) -> dict:
        table_path = os.path.join(self.table_home, table_name)
        if not os.path.exists(table_path):
            table_path = os.path.join(self.secondary_table_home, table_name)
            if not os.path.exists(table_path):
                raise FileNotFoundError(f'Can not find table {table_name} in {os.path.basename(self.table_home)} and {os.path.basename(self.secondary_table_home)}')
        return DefaultStaticTableLoader.excel_to_dict(table_path, primary_key_to_lower)

    @staticmethod
    def convert_table_value(value: str, value_type: str):
        """
        根据给定的类型将值进行转换。

        参数:
        - value: 要转换的值。
        - value_type: 目标数据类型。

        返回:
        - 转换后的值。
        """
        if value is None:
            value = ""

        try:
            if value_type == "str":
                return str(value)
            elif value_type == "int":
                return int(value) if value != "" else 0
            elif value_type == "float":
                return float(value) if value != "" else 0.0
            elif value_type == "bool":
                return str(value).lower() not in ["0", "false", ""]
            elif value_type == "list":
                return eval(value) if value != "" else []
            elif value_type == "dict":
                return eval(value) if value != "" else {}
            elif value_type == "json":
                return json.loads(value) if value != "" else None
            else:
                raise ValueError(f"Unsupported value type: {value_type}")
        except Exception as e:
            raise ValueError(f"Error converting value: {value} to {value_type}. {str(e)}")

    @staticmethod
    def excel_to_dict(file_path: str, primary_key_to_lower=False):
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook.worksheets[0]

        ret_dict = {}
        available_columns = []
        table_heads = {}

        for col_index in range(2, sheet.max_column + 1):
            cell_value = sheet.cell(row=2, column=col_index).value
            if cell_value in ["str", "int", "float", "bool", "list", "dict", "json", "/"]:
                if cell_value != "/":
                    available_columns.append((col_index, cell_value))
            elif cell_value is None or cell_value == "":
                break
            else:
                raise ValueError(f"Unsupported value type: {cell_value} at column {col_index}")

        for col_index, _ in available_columns:
            head_str = sheet.cell(row=1, column=col_index).value
            if head_str is None or head_str == "":
                raise ValueError(f"Empty header at column {col_index}")
            table_heads[col_index] = head_str

        empty_count = 0
        for row_index in range(4, sheet.max_row + 1):
            row_key = sheet.cell(row=row_index, column=2).value
            if row_key is None or row_key == "":
                empty_count += 1
                if empty_count >= 3:
                    break
                continue

            sub_dict = {'_i': len(ret_dict), '_ori_id': row_key}

            if primary_key_to_lower:
                row_key = row_key.lower()

            empty_count = 0

            for col_index, value_type in available_columns:
                cell_value = sheet.cell(row=row_index, column=col_index).value
                converted_value = DefaultStaticTableLoader.convert_table_value(cell_value, value_type)
                if col_index == 2 and primary_key_to_lower:
                    converted_value = converted_value.lower()
                sub_dict[table_heads[col_index]] = converted_value

            ret_dict[row_key] = sub_dict

        return ret_dict


if __name__ == "__main__":
    loader = DefaultStaticTableLoader()
    print(loader.load_table("table_set_group.xlsx"))
