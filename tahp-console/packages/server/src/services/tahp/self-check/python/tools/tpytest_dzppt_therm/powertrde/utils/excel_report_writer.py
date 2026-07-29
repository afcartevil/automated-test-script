from dataclasses import dataclass, field, fields, Field
from typing import Any, List, Optional, Tuple, Dict, Callable, Union
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from powertrde.utils.excel_utils import ExcelHyperLinkValue


def title_meta(
        *,
        default: Any = None,
        display_name: Union[None, str, Callable[[], str]] = None,
        display_name_ext: Optional[str] = None,
        width: Optional[float] = None,
        bg_color: Optional[str] = None,
        font_color: Optional[str] = None,
        header_align: Tuple[str, str] = ('center', 'center'),
        data_align: Tuple[str, str] = ('left', 'center'),
        number_format: Optional[str] = None,
        order: Optional[float] = None,
        is_template: bool = False
) -> Any:
    """
    Factory for dataclass fields to attach Excel header & data metadata.

    Non-title_meta fields will have empty metadata and be skipped in writing.

    :param default: default cell value
    :param display_name: header text (defaults to field name), or you can pass a text factory func
    :param display_name_ext: header text suffix
    :param width: column width
    :param bg_color: header background fill color
    :param font_color: header font color
    :param header_align: (horizontal, vertical) for header cell alignment
    :param data_align: (horizontal, vertical) for data cell alignment
    :param number_format: Excel number format string for data cells
    :param order: Sorting order of fields
    :param is_template: Specifies whether the field is a template

    Usage:
        @dataclass
        class Unit:
            col1: str = title_meta(
                default='', display_name='Column 1', width=20,
                bg_color='FFCCFF', font_color='000000',
                header_align=('center','center'), data_align=('left','center'),
                number_format=None
            )
    """
    return field(
        default=default,
        metadata={
            'display_name': display_name,
            'display_name_ext': display_name_ext,
            'width': width,
            'bg_color': bg_color,
            'font_color': font_color,
            'header_align': header_align,
            'data_align': data_align,
            'number_format': number_format,
            'order': order,
            'is_template': is_template
        }
    )


@dataclass
class DynamicField:
    filed_name: str
    field_value: Any
    field_desc: Field
    field_overwrites: Optional[dict] = None

    @classmethod
    def from_actual_field(cls, unit, actual_field_desc: Field):
        return DynamicField(filed_name=actual_field_desc.name, field_value=getattr(unit, actual_field_desc.name),
                            field_desc=actual_field_desc, field_overwrites=None)

    @classmethod
    def from_template(cls, filed_name: str, filed_value: Any, filed_desc: Field, **field_overwrites):
        return DynamicField(filed_name=filed_name, field_value=filed_value, field_desc=filed_desc,
                            field_overwrites=field_overwrites)

    def get_field_meta(self, meta_key: str, default_value=None):
        if self.field_overwrites and meta_key in self.field_overwrites:
            return self.field_overwrites[meta_key]

        m = self.field_desc.metadata
        if m:
            return m.get(meta_key, default_value)
        return default_value

    def get_field_order(self) -> Optional[int]:
        return self.get_field_meta('order')


@dataclass
class ExcelTableUnit:
    """
    Base class for Excel table units.
    Extend this class with your own fields, using title_meta to specify header & data settings.
    Non-title_meta fields are ignored when generating Excel output.
    Example:
        @dataclass
        class MyUnit(ExcelTableUnit):
            name: str = title_meta(
                default='', display_name='Name', width=25,
                bg_color='FFEB9C', font_color='000000',
                header_align=('center','center'), data_align=('left','center')
            )
            score: float = title_meta(
                default=0.0, display_name='Score', width=15,
                bg_color='C6EFCE', font_color='006100',
                number_format='0.00', data_align=('right','center')
            )
            remarks: str = title_meta(default='OK')
            # intermediate variable, not included in Excel:
            temp_cache: dict = field(default_factory=dict)
    """

    @classmethod
    def get_field_desc(cls, field_name: str) -> Optional[Field]:
        try:
            return getattr(cls, '__dataclass_fields__')[field_name]
        except Exception:
            return None

    def __post_init__(self):
        self.dynamic_fields: Dict[str, DynamicField] = {}

    def bind_dynamic_field(self, dynamic_field: DynamicField):
        self.dynamic_fields[dynamic_field.filed_name] = dynamic_field

    def _preprocess_source(self, source: dict):
        """
        将Source进行预处理，方便在后续parse中使用
        一般来说，你如果处理了source，应返回一个副本，防止直接影响来源
        """
        return source

    @classmethod
    def parse_from(cls, source: dict):
        ins = cls()
        ins.parse(source)
        return ins

    def parse(self, source: dict):
        _final_source = self._preprocess_source(source)
        for k, v in _final_source.items():
            if hasattr(self, k):
                if not self._on_parse_field(k, v):
                    setattr(self, k, v)
        self._on_parse_finished(_final_source)

    def _on_parse_field(self, filed_name: str, source_value: Any) -> bool:
        return False

    def _on_parse_finished(self, final_source: dict):
        pass


class ExcelTableWriter:
    """
    Write lists of dataclass-based units to sheets in an Excel workbook.
    """

    def __init__(self, excel_filepath: str):
        self.excel_filepath = excel_filepath
        self.workbook = self._create_or_load_workbook(excel_filepath)
        # default styles
        self._default_fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
        self._default_font = Font(name='微软雅黑', bold=True)
        self._default_border = Border(*(Side(style='thin'),) * 4)

    @staticmethod
    def _create_or_load_workbook(path: str) -> openpyxl.Workbook:
        try:
            # 如果传入的是目录或不可读对象，则直接新建
            import os
            if (not path) or (os.path.isdir(path)):
                wb = openpyxl.Workbook()
                try:
                    wb.remove(wb.active)
                except Exception:
                    pass
                return wb
            # 尝试加载已有工作簿
            wb = openpyxl.load_workbook(path)
            return wb
        except Exception:
            # 加载失败（损坏/非xlsx/占位文件等），回退为新建工作簿
            wb = openpyxl.Workbook()
            try:
                wb.remove(wb.active)
            except Exception:
                pass
            return wb

    def save(self) -> None:
        """Save workbook to disk."""
        self.workbook.save(self.excel_filepath)

    @staticmethod
    def to_unit_dynamic_fields_pairs(units: List[ExcelTableUnit]) -> List[Tuple[ExcelTableUnit, Dict[str, DynamicField]]]:
        ret = []
        for u in units:
            a = {f.name: DynamicField.from_actual_field(u, f) for f in fields(u) if (f.metadata and not f.metadata.get('is_template'))}
            if u.dynamic_fields:
                a.update(u.dynamic_fields)
            ret.append((u, a))
        return ret

    @staticmethod
    def _to_sorted_dynamic_fields(pairs: List[Tuple[ExcelTableUnit, Dict[str, DynamicField]]]):
        _no_order_fields: Dict[str, float] = {}
        _all_fields: Dict[str, DynamicField] = {}
        for t in pairs:
            for d in t[1].values():
                if d.filed_name not in _all_fields:
                    _all_fields[d.filed_name] = d

        _all_fields_with_order = list(_all_fields.values())

        for df in _all_fields_with_order:
            df_order = df.get_field_order()
            if df_order is None:
                _no_order_fields[df.filed_name] = float(len(_no_order_fields))

        def _get_order(_df: DynamicField) -> float:
            _o = _no_order_fields.get(_df.filed_name, None)
            if _o is None:
                return _df.get_field_order()
            return _o

        _all_fields_with_order.sort(key=lambda _df: _get_order(_df))
        return _all_fields_with_order

    def write_sheet(self, sheet_name: str, units: List[ExcelTableUnit], save_workbook=True) -> None:
        # remove existing sheet
        if sheet_name in self.workbook.sheetnames:
            self.workbook.remove(self.workbook[sheet_name])
        if not units:
            if save_workbook:
                self.save()
            return

        ws = self.workbook.create_sheet(title=sheet_name)
        # sample = units[0]
        cols = []  # (name, metadata...)

        all_unit_df_pairs = self.to_unit_dynamic_fields_pairs(units)
        sample_sorted_fields = self._to_sorted_dynamic_fields(all_unit_df_pairs)
        max_header_lines = 1

        for f in sample_sorted_fields:
            # # skip fields without title_meta metadata
            # md = f.metadata
            # if not md:
            #     continue
            header = f.get_field_meta('display_name') or f.filed_name
            if callable(header):
                header = header()
            header_ext = f.get_field_meta('display_name_ext')
            if header_ext:
                header += header_ext
            width = f.get_field_meta('width')
            bg = f.get_field_meta('bg_color')
            fg = f.get_field_meta('font_color')
            head_align = f.get_field_meta('header_align', ('center', 'center'))
            data_align = f.get_field_meta('data_align', ('left', 'center'))
            num_fmt = f.get_field_meta('number_format')
            # 如果未指定宽度，则根据字符长度推断（中文按2，其他按1），再加2
            # 如果指定了宽度，则作为最小宽度，实际宽度取计算值和指定值中的较大者
            # 考虑换行符，找出最长的一行来计算宽度
            max_line_length = 0
            header_lines = header.split('\n')
            max_header_lines = max(max_header_lines, len(header_lines))
            for line in header_lines:
                # 处理\r\n情况
                line = line.rstrip('\r')
                line_length = 0
                for ch in line:
                    line_length += 2 if '\u4e00' <= ch <= '\u9fff' else 1
                max_line_length = max(max_line_length, line_length)
            calculated_width = max_line_length + 2
            
            # 如果用户指定了宽度，取计算值和指定值中的较大者
            if width is not None:
                width = max(width, calculated_width)
            else:
                width = calculated_width
            cols.append((f.filed_name, header, width, bg, fg, head_align, data_align, num_fmt))

        # write header
        for idx, (_, header, width, bg, fg, head_align, _, _) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=idx, value=header)
            if width:
                ws.column_dimensions[cell.column_letter].width = width
            # apply header styles
            cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type='solid') if bg else self._default_fill
            cell.font = Font(name='微软雅黑', bold=True, color=fg) if fg else self._default_font
            cell.alignment = Alignment(horizontal=head_align[0], vertical=head_align[1], wrap_text=('\n' in header))
            cell.border = self._default_border

        # header row height & freeze
        base_height = ws.row_dimensions[1].height or 15
        ws.row_dimensions[1].height = base_height * max_header_lines + 0.5 * base_height
        ws.freeze_panes = 'A2'

        # write data rows
        for r, unit_pair in enumerate(all_unit_df_pairs, start=2):
            for c, (field_name, *_, data_align, num_fmt) in enumerate(cols, start=1):
                df = unit_pair[1].get(field_name)
                raw_value = None if df is None else df.field_value
                if isinstance(raw_value, ExcelHyperLinkValue):
                    cell = ws.cell(row=r, column=c, value=raw_value.value)
                    cell.hyperlink = raw_value.link
                    try:
                        cell.style = "Hyperlink"
                    except Exception:
                        pass
                    cell.alignment = Alignment(horizontal=data_align[0], vertical=data_align[1], wrap_text=(isinstance(raw_value.value, str) and '\n' in str(raw_value.value)))
                    if num_fmt:
                        cell.number_format = num_fmt
                    cell.border = self._default_border
                else:
                    df_value = raw_value
                    if isinstance(df_value, bool):
                        df_value = 1 if df_value else 0
                    elif not isinstance(df_value, (int, float, str)):
                        df_value = str(df_value)
                    cell = ws.cell(row=r, column=c, value=df_value)
                    cell.alignment = Alignment(horizontal=data_align[0], vertical=data_align[1], wrap_text=(isinstance(df_value, str) and '\n' in df_value))
                    if num_fmt:
                        cell.number_format = num_fmt
                    cell.border = self._default_border

        if save_workbook:
            self.save()
