import os
import re
from collections import defaultdict
from openpyxl import Workbook

# NOTE: SOURCE_FILES / CONFIG_FILES point to local Android source trees.
# Update these paths on each host before running refresh_matrix.py.
# Output is always written to ../data/carrier_keys_matrix.xlsx (portable).

SOURCE_FILES = [
    ("mtk", r"Z:\mtk_w\w\vendor\mediatek\proprietary\frameworks\opt\telephony-base\java\mediatek\telephony\MtkCarrierConfigManager.java", "MtkCarrierConfigManager"),
    ("mtk", r"Z:\mtk_w\w\frameworks\base\telephony\java\android\telephony\CarrierConfigManager.java", "CarrierConfigManager"),
    ("sprd", r"Z:\sprd_w\sys\frameworks\base\telephony\java\android\telephony\CarrierConfigManager.java", "CarrierConfigManager"),
    ("sprd", r"Z:\sprd_w\sys\vendor\sprd\ex-interface\frameworks\base\telephony\java\android\telephony\UniCarrierConfigManager.java", "UniCarrierConfigManager"),
    ("qcom", r"Z:\q_w\system\vendor\codeaurora\commonsys\telephony\ims\ims-ext-common\src\org\codeaurora\ims\QtiCarrierConfigs.java", "QtiCarrierConfigs"),
    ("qcom", r"Z:\q_w\system\frameworks\base\telephony\java\android\telephony\CarrierConfigManager.java", "CarrierConfigManager"),
]

CONFIG_FILES = {
    "mtk": r"Z:\mtk_w\w\vendor\transsion\system_modules\packages\apps\trancarrierconfig\trancarrierconfig\carrier_config_mtk_all_project.xml",
    "sprd": r"Z:\mtk_w\w\vendor\transsion\system_modules\packages\apps\trancarrierconfig\trancarrierconfig\carrier_config_sprd_all_project.xml",
    "qcom": r"Z:\mtk_w\w\vendor\transsion\system_modules\packages\apps\trancarrierconfig\trancarrierconfig\carrier_config_qcom_all_project.xml",
}

COMMON_GLOB_PREFIXES = ("carrier_config_mccmnc_", "carrier_config_carrierid_")
CARRIER_LIST_FILE = r"Z:\mtk_w\w\packages\providers\TelephonyProvider\assets\latest_carrier_id\carrier_list.textpb"

DEFAULT_LITERAL_RE = re.compile(r'(?:sDefaults|defaults|defaultValues)\.put[A-Za-z]+\(\s*"([^"]+)"\s*,\s*(.+?)\s*\);')
DEFAULT_CONST_RE = re.compile(r'(?:sDefaults|defaults|defaultValues)\.put[A-Za-z]+\(\s*([A-Z0-9_]+)\s*,\s*(.+?)\s*\);')

def infer_comment(lines, idx):
    comments = []
    i = idx - 1
    while i >= 0:
        s = lines[i].strip()
        if not s:
            i -= 1
            continue
        if s.startswith("@"):
            i -= 1
            continue
        if s.startswith("//"):
            comments.append(s.lstrip("/ ").strip())
            i -= 1
            continue
        if s.endswith("*/"):
            block = []
            s_stripped = s.rstrip("*/").strip()
            if s_stripped and not s_stripped.startswith("/*"):
                block.append(s_stripped.lstrip("* ").strip())
            i -= 1
            while i >= 0:
                bs = lines[i].strip()
                is_start = bs.startswith("/**") or bs.startswith("/*")
                if is_start:
                    content = bs.replace("/**", "").replace("/*", "").strip()
                    if content:
                        block.append(content.lstrip("* ").strip())
                    break
                else:
                    block.append(bs.lstrip("* ").strip())
                i -= 1
            block.reverse()
            comments.extend([x for x in block if x])
            i -= 1
            continue
        break
    comments.reverse()
    return " ".join(comments)

def parse_java_file(platform, path, source_name):
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    lines = content.splitlines()

    line_starts = []
    curr = 0
    for line in lines:
        line_starts.append(curr)
        curr += len(line) + 1

    def get_line_idx(pos):
        for i, start in enumerate(line_starts):
            if start > pos:
                return max(0, i - 1)
        return len(lines) - 1

    const_pattern = re.compile(r"(?:public|private|protected)?\s*static\s+final\s+String\s+([A-Z0-9_]+)\s*=\s*(?:([A-Z0-9_]+)\s*\+\s*)?\"([^\"]+)\"", re.DOTALL)
    
    string_vars = {}
    const_map = {}
    const_info = {}
    for m in const_pattern.finditer(content):
        const_name = m.group(1)
        prefix_var = m.group(2)
        val = m.group(3)

        if not prefix_var:
            string_vars[const_name] = val
            key_name = val
        else:
            key_name = string_vars.get(prefix_var, "") + val
            string_vars[const_name] = key_name

        if const_name in ("KEY_PREFIX", "TAG"):
            continue

        line_idx = get_line_idx(m.start())
        feature_desc = infer_comment(lines, line_idx)

        const_map[const_name] = key_name
        const_info[key_name] = {
            "platform": platform,
            "source_file": path,
            "source_class": source_name,
            "const_name": const_name,
            "key_name": key_name,
            "feature_desc": feature_desc,
            "default_value": "",
        }

    defaults = {}
    for line in lines:
        m = DEFAULT_LITERAL_RE.search(line)
        if m:
            defaults[m.group(1)] = m.group(2).strip()
            continue
        mc = DEFAULT_CONST_RE.search(line)
        if mc:
            key_const = mc.group(1)
            key_name = const_map.get(key_const, key_const)
            defaults[key_name] = mc.group(2).strip()

    for key_name, dv in defaults.items():
        if key_name in const_info:
            const_info[key_name]["default_value"] = dv
        else:
            const_info[key_name] = {
                "platform": platform,
                "source_file": path,
                "source_class": source_name,
                "const_name": "",
                "key_name": key_name,
                "feature_desc": "",
                "default_value": dv,
            }
    return list(const_info.values())

def parse_existing_all_project_keys(path):
    keys = set()
    if not os.path.exists(path):
        return keys
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            m = re.search(r'name="([^"]+)"', line)
            if m:
                keys.add(m.group(1))
    return keys

def collect_common_support(base_dir):
    common_keys = set()
    if not os.path.exists(base_dir):
        return common_keys
    for root, _, files in os.walk(base_dir):
        for fn in files:
            if fn.startswith(COMMON_GLOB_PREFIXES) and fn.endswith(".xml"):
                fp = os.path.join(root, fn)
                with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                    for line in f:
                        m = re.search(r'name="([^"]+)"', line)
                        if m:
                            common_keys.add(m.group(1))
    return common_keys

def normalize_feature(row):
    desc = row["feature_desc"].strip()
    if desc:
        return desc
    const_name = row["const_name"].lower()
    return const_name

def looks_like_carrier_key(row):
    const_name = row["const_name"]
    key_name = row["key_name"]
    if const_name.startswith("KEY_") or const_name.startswith("MTK_KEY_") or const_name.startswith("TRAN_KEY_"):
        return True
    if key_name.startswith("carrier_") or key_name.startswith("mtk_") or key_name.startswith("tran_") or key_name.startswith("apn."):
        return True
    if "_" in key_name and key_name.lower() == key_name:
        return True
    if "." in key_name and key_name.lower() == key_name:
        return True
    return False

def build_xml_only_rows(rows, all_project_keys):
    existing = {(r["platform"], r["key_name"]) for r in rows}
    xml_rows = []
    for platform, keys in all_project_keys.items():
        for key_name in sorted(keys):
            if (platform, key_name) in existing:
                continue
            xml_rows.append({
                "platform": platform,
                "source_file": CONFIG_FILES[platform],
                "source_class": "all_project_xml",
                "const_name": "",
                "key_name": key_name,
                "feature_desc": "from platform all_project xml",
                "default_value": "",
            })
    return xml_rows

def parse_textpb_scalar(line):
    key, value = [x.strip() for x in line.split(":", 1)]
    if value.startswith('"') and value.endswith('"'):
        value = value[1:-1]
    return key, value

def parse_carrier_list_textpb(path):
    carriers = []
    if not os.path.exists(path):
        return carriers
    current_carrier = None
    current_attr = None

    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for raw in f:
            line = raw.strip()
            if not line:
                continue

            if line == "carrier_id {":
                current_carrier = {
                    "canonical_id": "",
                    "carrier_name": "",
                    "top_level_fields": defaultdict(list),
                    "attributes": [],
                }
                current_attr = None
                continue

            if line == "carrier_attribute {":
                current_attr = defaultdict(list)
                continue

            if line == "}":
                if current_attr is not None and current_carrier is not None:
                    current_carrier["attributes"].append(dict(current_attr))
                    current_attr = None
                elif current_carrier is not None:
                    carriers.append(current_carrier)
                    current_carrier = None
                continue

            if ":" not in line or current_carrier is None:
                continue

            key, value = parse_textpb_scalar(line)
            if current_attr is not None:
                current_attr[key].append(value)
            else:
                if key == "canonical_id":
                    current_carrier["canonical_id"] = value
                elif key == "carrier_name":
                    current_carrier["carrier_name"] = value
                else:
                    current_carrier["top_level_fields"][key].append(value)

    return carriers

def build_carrier_list_rows(carriers):
    rows = []
    mvno_keys = [
        "gid1",
        "gid2",
        "spn",
        "imsi_prefix_xpattern",
        "iccid_prefix",
        "privilege_access_rule",
    ]
    extra_keys = set()
    for carrier in carriers:
        for attr in carrier["attributes"]:
            extra_keys.update(attr.keys())

    preferred_order = [
        "mccmnc_tuple",
        "gid1",
        "gid2",
        "spn",
        "imsi_prefix_xpattern",
        "iccid_prefix",
        "privilege_access_rule",
        "apn",
        "plmn",
    ]
    extra_keys = [k for k in preferred_order if k in extra_keys] + sorted(
        [k for k in extra_keys if k not in preferred_order]
    )

    for carrier in carriers:
        cid = carrier["canonical_id"]
        cname = carrier["carrier_name"]
        safe_name = re.sub(r"[^0-9A-Za-z_-]", "-", cname).replace(" ", "-")
        safe_name = re.sub(r"-{2,}", "-", safe_name).strip("-")
        top_fields = {
            k: "|".join(v) if isinstance(v, list) else str(v)
            for k, v in carrier["top_level_fields"].items()
        }

        if not carrier["attributes"]:
            row = {
                "canonical_id": cid,
                "carrier_name": cname,
                "attribute_index": 0,
                "is_mvno": "N",
                "mvno_match_type": "",
                "suggested_file_name": f"carrier_config_carrierid_{cid}_{safe_name}.xml",
            }
            row.update(top_fields)
            rows.append(row)
            continue

        for idx, attr in enumerate(carrier["attributes"], start=1):
            present_mvno = [k for k in mvno_keys if attr.get(k)]
            mno_or_mvno = "虚拟运营商(MVNO)" if present_mvno else "基础运营商(MNO)"
            match_parts = []
            for k in extra_keys:
                vals = attr.get(k, [])
                if vals:
                    match_parts.append(f"{k}=" + "|".join(vals))
            row = {
                "canonical_id": cid,
                "carrier_name": cname,
                "attribute_index": idx,
                "mno_or_mvno": mno_or_mvno,
                "is_mvno": "Y" if present_mvno else "N",
                "mvno_match_type": "|".join(present_mvno),
                "match_rule_summary": "; ".join(match_parts),
                "suggested_file_name": f"carrier_config_carrierid_{cid}_{safe_name}.xml",
            }
            row.update(top_fields)
            for k in extra_keys:
                row[k] = "|".join(attr.get(k, []))
            rows.append(row)

    rows.sort(key=lambda x: (int(x.get("canonical_id", 0) or 0), int(x.get("attribute_index", 0) or 0)))
    return rows

import openpyxl

def write_excel(rows, carrier_records, out_file):
    if os.path.exists(out_file):
        try:
            wb = openpyxl.load_workbook(out_file)
        except Exception:
            wb = Workbook()
    else:
        wb = Workbook()

    if "carrier_keys" in wb.sheetnames:
        del wb["carrier_keys"]
    
    if "Carrier_Config" in wb.sheetnames:
        del wb["Carrier_Config"]

    ws1 = wb.create_sheet("carrier_keys", 0)  # Insert as first sheet
    
    fields = [
        "supported_platforms",
        "const_name",
        "key_name",
        "feature_desc",
        "default_value",
    ]
    ws1.append(fields)

    def excel_safe(v):
        s = "" if v is None else str(v)
        if s.startswith("="):
            return "'" + s
        return s

    for row in rows:
        ws1.append([excel_safe(row.get(k, "")) for k in fields])

    wb.save(out_file)

def resolve_output_dir():
    """Output next to shared/data/ so the script stays portable with the skills bundle."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.normpath(os.path.join(script_dir, "..", "data"))
    os.makedirs(out_dir, exist_ok=True)
    return out_dir


def main():
    out_dir = resolve_output_dir()

    rows = []
    for platform, path, source_name in SOURCE_FILES:
        rows.extend(parse_java_file(platform, path, source_name))

    rows = [r for r in rows if looks_like_carrier_key(r)]

    all_project_keys = {p: parse_existing_all_project_keys(fp) for p, fp in CONFIG_FILES.items()}
    rows.extend(build_xml_only_rows(rows, all_project_keys))
    
    common_xml_keys = collect_common_support(r"Z:\mtk_w\w\vendor\transsion\system_modules\packages\apps\trancarrierconfig\trancarrierconfig")

    key_platforms = defaultdict(set)
    for row in rows:
        key_platforms[row["key_name"]].add(row["platform"])

    merged_rows = {}
    for row in rows:
        kn = row["key_name"]
        if kn not in merged_rows:
            merged_rows[kn] = row.copy()
        else:
            old = merged_rows[kn]
            if not old["default_value"] and row.get("default_value"):
                old["default_value"] = row["default_value"]
            if not old["feature_desc"] and row.get("feature_desc"):
                old["feature_desc"] = row["feature_desc"]
            if not old["const_name"] and row.get("const_name"):
                old["const_name"] = row["const_name"]
            
            sf = old["source_file"]
            if row["source_file"] not in sf:
                old["source_file"] = sf + " | " + row["source_file"]
                
            sc = old["source_class"]
            if row["source_class"] not in sc:
                old["source_class"] = sc + " | " + row["source_class"]

    final_rows = []
    summary = defaultdict(int)
    
    # We enforce specific platform ordering: MTK SPRD QCOM
    platform_order = {"mtk": 1, "sprd": 2, "qcom": 3}
    
    for kn, row in merged_rows.items():
        plats = key_platforms[kn]
        ordered_plats = sorted(list(plats), key=lambda x: platform_order.get(x, 99))
        supported_platforms_str = " ".join([p.upper() for p in ordered_plats])
        
        row["supported_platforms"] = supported_platforms_str
        row["platform"] = supported_platforms_str
        
        if len(plats) == 3:
            row["recommend_target"] = "common_xml(carrier_config_mccmnc_*/carrier_config_carrierid_*)"
        else:
            # We can just join them
            p_str = "_".join(ordered_plats)
            row["recommend_target"] = f"{p_str}_all_project"
            
        row["already_in_common_xml"] = "Y" if kn in common_xml_keys else "N"
        
        already_in_plat = "N"
        for p in plats:
            if kn in all_project_keys[p]:
                already_in_plat = "Y"
                break
        row["already_in_platform_all_project"] = already_in_plat
        
        row["feature_desc"] = normalize_feature(row)
        final_rows.append(row)
        
        for p in plats:
            summary[p] += 1

    final_rows.sort(key=lambda x: (x["platform"], x["source_class"], x["key_name"]))

    carrier_records = parse_carrier_list_textpb(CARRIER_LIST_FILE)
    xlsx_file = os.path.join(out_dir, "carrier_keys_matrix.xlsx")
    write_excel(final_rows, carrier_records, xlsx_file)

    print(f"mtk: {summary['mtk']} keys")
    print(f"sprd: {summary['sprd']} keys")
    print(f"qcom: {summary['qcom']} keys")
    print(f"total unique keys: {len(final_rows)}")
    print(xlsx_file)

if __name__ == "__main__":
    main()